#!/usr/bin/env python3
"""
gen_eecc_v9.py — Generador genérico de Estados Contables (RT 54 / RP CPCECABA 460/2024)

Lee:
  * el Sumas y Saldos del ejercicio (Excel exportado del sistema: Cuenta | Saldo,
    sin columna de rubro, antes del cierre y en valores históricos),
  * el EECC del ejercicio anterior (PDF) para comparativos, PN de inicio y Anexo I.

Calcula (simplificaciones para entidades pequeñas, RT 54):
  * PN de inicio reexpresado con el coeficiente,
  * amortizaciones del ejercicio (Anexo I anterior + altas, año de alta completo),
  * RECPAM / resultados financieros y por tenencia por diferencia.

Genera un .xlsx (con hoja "Control" de clasificación) y un .pdf con:
  Carátula, ESP, ER, EEPN, EFE (método indirecto), Anexos y Notas.

Si algo no se puede resolver (cuenta desconocida, SS que no cuadra) termina con
código 2 y el motivo en stderr, para que el servidor marque "Error" en Notion.
"""
import argparse
import json
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta

# ════════════════════════════════════════════════════════════════════════════
# Utilidades
# ════════════════════════════════════════════════════════════════════════════


class EECCError(Exception):
    pass


def _strip_char(c):
    d = unicodedata.normalize("NFKD", c)
    base = "".join(ch for ch in d if not unicodedata.combining(ch))
    if len(base) != 1:
        base = c
    return base.lower()


def norm(s):
    """minúsculas, sin acentos; conserva la longitud del texto."""
    return "".join(_strip_char(c) for c in str(s))


def skey(s):
    """clave tolerante a plurales: cada palabra sin 's' final."""
    return " ".join(w[:-1] if len(w) > 3 and w.endswith("s") else w for w in nkey(s).split())


def nkey(s):
    """clave de comparación: norm + espacios colapsados, sin $ ni puntuación final."""
    s = norm(s).replace("$", " ").replace("º", "o").replace("°", "o")
    s = re.sub(r"[\s.:]+$", "", s)
    return re.sub(r"\s+", " ", s).strip()


NUM_RE = re.compile(
    r'\(\s*-?\d{1,3}(?:[.,]\d{3})*[.,]\d{2}\s*\)'
    r'|-?\d{1,3}(?:[.,]\d{3})+[.,]\d{2}'
    r'|-?\d+[.,]\d{2}(?!\d)'
    r'|(?<![\w.,])"?[-–]"?(?![\w\d])'
)


def parse_num(tok):
    t = tok.strip().strip('"')
    if t in ("-", "–", ""):
        return 0.0
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()").strip()
    if t.startswith("-"):
        neg, t = True, t[1:]
    m = re.match(r"^([\d.,]*?)([.,])(\d{2})$", t)
    if not m:
        return None
    ent = re.sub(r"[.,]", "", m.group(1))
    v = float(f"{ent or 0}.{m.group(3)}")
    return -v if neg else v


def tokens_after(line, pos):
    out = []
    for m in NUM_RE.finditer(line, pos):
        v = parse_num(m.group())
        if v is not None:
            out.append(v)
    return out


def fmt_ar(v, dash=True):
    if v is None:
        return ""
    if abs(v) < 0.005:
        return "-" if dash else "0,00"
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"({s})" if v < 0 else s


def fmt_cuit(c):
    d = re.sub(r"\D", "", str(c))
    return f"{d[:2]}-{d[2:10]}-{d[10:]}" if len(d) == 11 else str(c)


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def fecha_larga(d):
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


# ════════════════════════════════════════════════════════════════════════════
# Clasificación de cuentas
# ════════════════════════════════════════════════════════════════════════════
RUBROS = {
    "caja": ("A", "Caja y bancos"),
    "inv": ("A", "Inversiones"),
    "cv": ("A", "Créditos por ventas"),
    "oc": ("A", "Otros créditos"),
    "bc": ("A", "Bienes de cambio"),
    "bu": ("A", "Bienes de uso"),
    "dc": ("P", "Comerciales"),
    "prest": ("P", "Préstamos"),
    "rem": ("P", "Remuneraciones y cargas sociales"),
    "df": ("P", "Cargas fiscales"),
    "ds": ("P", "Deudas con socios"),
    "od": ("P", "Otras deudas"),
    "cap": ("PN", "Capital social"),
    "aj": ("PN", "Ajuste de capital"),
    "apirr": ("PN", "Aportes irrevocables"),
    "resleg": ("PN", "Reservas"),
    "rna": ("PN", "Resultados no asignados"),
    "ventas": ("R", "Ventas"),
    "costo": ("R", "Costo"),
    "gcom": ("R", "Gastos de comercialización"),
    "gadm": ("R", "Gastos de administración"),
    "amort": ("R", "Amortizaciones"),
    "rfyt": ("R", "Resultados financieros y por tenencia"),
    "oi": ("R", "Otros ingresos y egresos"),
    "ig": ("R", "Impuesto a las ganancias"),
}

RULES = [
    # Patrimonio neto
    ("aj", r"ajuste (de|del) capital"),
    ("apirr", r"aportes? irrevocables?"),
    ("resleg", r"^reserva"),
    ("cap", r"^capital( social| suscripto| integrado)?\b"),
    ("rna", r"resultados? no asignados?|^rna$|resultados? del ejercicio|resultados? acumulados?|"
            r"resultados? (de )?ejercicios? anteriores?"),
    # Resultados financieros (se recalculan por diferencia)
    ("rfyt", r"recpam|diferencias? de cambio|redondeo|resultados? por tenencia|resultados? financieros|"
             r"^intereses|intereses (ganados|perdidos|cobrados|pagados|bancarios)"),
    # Bienes de uso y amortizaciones
    ("bu", r"amortizacion(es)? acumulada|amort\.? acum|depreciacion acumulada"),
    ("amort", r"amortizacion|depreciacion"),
    ("bu", r"muebles|\butiles\b|rodados?$|^rodado|instalaciones|equipos?\b|maquinari|herramientas|"
           r"inmueble|edificio|terreno|computacion|computadoras|^mejoras|hardware"),
    # Pasivo (antes que resultados: "a pagar")
    ("rem", r"(sueldos?|jornales|cargas sociales|sindicato|seguro de vida|sac\b|aguinaldo|vacaciones|"
            r"\bart\b|obra social|f\.? ?931|aportes y contribuciones).*(a pagar|devengad)|"
            r"remuneraciones a pagar"),
    ("df", r"(iva|iibb|ingresos brutos|ganancias|impuestos?|tasas?|abl|monotributo|bienes personales|"
           r"sicore|sircreb|percepcion|retencion).*a pagar|debito fiscal|"
           r"(retencion|retenciones|percepcion|percepciones).*efectuad|"
           r"provision (de |del )?(ig|impuesto|ganancias)|plan(es)? de pagos?|moratoria"),
    ("ds", r"deuda.*socio|socio.*a pagar|cuenta particular|cuentas? (corriente )?socios?|"
           r"dividendos a pagar|honorarios (de )?directores a pagar|aportes? a reintegrar"),
    ("prest", r"prestamo|descubierto|adelanto en cuenta"),
    ("dc", r"proveedor|acreedores comerciales|documentos a pagar|cheques diferidos|tarjetas? de credito"),
    ("od", r"acreedores varios|otras deudas|anticipos? de clientes|deudas varias"),
    # Activo
    ("caja", r"^caja|banco|mercado ?pago|fondo fijo|valores a depositar|moneda extranjera|^efectivo"),
    ("inv", r"plazo fijo|fondos? comun|fondos? de inversion|^fci|^inversion|cauciones?|^acciones|^bonos"),
    ("cv", r"deudores por ventas?|^clientes|documentos a cobrar|cheques en cartera|tarjetas? a cobrar|"
           r"creditos por ventas?"),
    ("bc", r"mercaderias?$|^mercaderia|bienes de cambio|materias primas|productos (en proceso|terminados)|"
           r"^stock|existencias?"),
    ("oc", r"iva (saldo|credito|a favor)|saldo (tecnico|libre)|libre disponibilidad|^saf\b|saldo a favor|"
           r"(retencion|retenciones|percepcion|percepciones).*(sufrid|iibb|iva|ganancias)|"
           r"anticipos?|pagos? a cuenta|25\.?413|\b(db|debitos?) y (cr|creditos?)|credito fiscal|"
           r"^socio|gastos pagados por adelantado|creditos? (varios|diversos)|deudores varios|"
           r"prestamos? a socios|a devengar|depositos? en garantia|reintegros? a cobrar|"
           r"integracion (de )?capital|capital (suscripto )?(pendiente|a integrar)|suscripcion pendiente"),
    # Resultados
    ("costo", r"^compras?\b|compra de|^costo|\bcmv\b|mercaderias vendidas"),
    ("ventas", r"^ventas?\b|venta de|ingresos? por|servicios prestados|honorarios (ganados|cobrados|percibidos)|"
               r"^facturacion|ingresos? (operativos|de explotacion)"),
    ("ig", r"impuesto a las ganancias"),
    ("oi", r"otros ingresos|otros egresos|resultado (por )?venta (de )?bienes de uso|recupero|bonifi"),
    ("gcom", r"iibb|ingresos brutos|publicidad|propaganda|comercializacion|fletes|comisiones|marketing|"
             r"convenio multilateral"),
    ("gadm", r"gasto|sueldo|jornal|cargas sociales|honorario|servicios publicos|alquiler|mantenimiento|"
             r"libreria|papeleria|seguro|impuesto|tasa|movilidad|viatico|combustible|telefon|internet|"
             r"\bluz\b|\bagua\b|\bgas\b|abl|monotributo|sistema|software|donacion|multa|refrigerio|"
             r"limpieza|correo|vigilancia|capacitacion|insumos|reparacion|expensas|bancari|comision|"
             r"cuota|suscripcion|indemnizacion|obra social|contribucion|aportes|remuneracion|premio"),
]
RULES_C = [(k, re.compile(p)) for k, p in RULES]


def clasificar(nombre):
    n = nkey(nombre)
    for k, rx in RULES_C:
        if rx.search(n):
            return k
    return None


MONEDA_EXT = re.compile(r"\busd\b|u\$s|dolar|euro|moneda extranjera|\bme\b")

BU_TIPOS = [
    ("Muebles y útiles", r"mueble|utiles", 10.0),
    ("Rodados", r"rodado", 20.0),
    ("Instalaciones", r"instalacion", 10.0),
    ("Equipos de computación", r"computacion|computadora|hardware", 33.33),
    ("Equipos y herramientas", r"equipo|herramienta|maquinari", 10.0),
    ("Inmueble terreno", r"terreno", 0.0),
    ("Inmueble edificio", r"inmueble|edificio|mejoras", 2.0),
]


def tipo_bu(nombre):
    n = nkey(nombre)
    for t, p, _ in BU_TIPOS:
        if re.search(p, n):
            return t
    return "Otros bienes de uso"


def tasa_default(tipo):
    for t, _, r in BU_TIPOS:
        if t == tipo:
            return r
    return 10.0


# ════════════════════════════════════════════════════════════════════════════
# Lectura del Sumas y Saldos
# ════════════════════════════════════════════════════════════════════════════
def leer_ss(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # si la hoja activa está vacía, usar la primera hoja con datos
    if not any(any(c is not None for c in r) for r in rows):
        for w in wb.worksheets:
            rr = [list(r) for r in w.iter_rows(values_only=True)]
            if any(any(c is not None for c in r) for r in rr):
                rows = rr
                break

    hdr_i, c_cta, c_sal, c_debe, c_haber = None, None, None, None, None
    c_hist, solo_ajustado = None, False
    hdr_txt = ""
    for i, r in enumerate(rows[:30]):
        labels = [nkey(c) if isinstance(c, str) else "" for c in r]
        if any(l.startswith("cuenta") or l in ("descripcion", "nombre", "denominacion", "nombre de la cuenta")
               for l in labels):
            hdr_i = i
            hdr_txt = " | ".join(str(c) for c in r if c is not None)
            sal_hist, sal_any, sal_adj = None, None, None
            for j, l in enumerate(labels):
                if c_cta is None and (l.startswith("cuenta") or l in ("descripcion", "nombre", "denominacion",
                                                                       "nombre de la cuenta")):
                    c_cta = j
                if "saldo" in l and "anterior" not in l:
                    sal_any = j
                    if "ajustad" in l:
                        sal_adj = j          # saldo ajustado por inflación (Xubio)
                    else:
                        sal_hist = j         # saldo histórico (de libros)
                if (l.startswith("debe") or l.startswith("debito")) and "ajustad" not in l:
                    c_debe = j
                if (l.startswith("haber") or l.startswith("credito")) and "ajustad" not in l:
                    c_haber = j
            # los estados se arman con el saldo ajustado si está; el histórico se usa para Ganancias y asientos
            c_sal = sal_adj if sal_adj is not None else (sal_hist if sal_hist is not None else sal_any)
            c_hist = sal_hist if (sal_adj is not None and sal_hist is not None) else None
            solo_ajustado = sal_adj is not None and sal_hist is None
            break
    if hdr_i is None:
        # sin encabezado: primera columna de texto + última numérica
        hdr_i = -1
        for r in rows:
            txt = [j for j, c in enumerate(r) if isinstance(c, str) and c.strip()]
            num = [j for j, c in enumerate(r) if isinstance(c, (int, float))]
            if txt and num:
                c_cta, c_sal = txt[0], num[-1]
                break
    if c_cta is None or (c_sal is None and (c_debe is None or c_haber is None)):
        raise EECCError("No pude identificar las columnas Cuenta / Saldo en el sumas y saldos. "
                        + (f"Encabezado encontrado: {hdr_txt}" if hdr_txt else
                           "No encontré una fila de encabezado con la palabra «Cuenta».")
                        + " El archivo tiene que tener una columna «Cuenta» y una «Saldo» (o Debe/Haber).")

    cuentas, hist = [], []
    for r in rows[hdr_i + 1:]:
        if c_cta >= len(r):
            continue
        nom = r[c_cta]
        if not isinstance(nom, str) or not nom.strip():
            continue
        if re.match(r"^(total|totales|sumas)\b", nkey(nom)):
            continue
        if c_sal is not None:
            v = r[c_sal] if c_sal < len(r) else None
        else:
            d = r[c_debe] or 0
            h = r[c_haber] or 0
            d = d if isinstance(d, (int, float)) else 0
            h = h if isinstance(h, (int, float)) else 0
            v = d + h if h < 0 else d - h      # Xubio exporta los créditos con signo negativo
        if isinstance(v, str):
            v = parse_num(v)
        if v is None:
            v = 0.0
        cuentas.append((nom.strip(), float(v)))
        if c_hist is not None:
            vh = r[c_hist] if c_hist < len(r) else None
            vh = parse_num(vh) if isinstance(vh, str) else vh
            hist.append((nom.strip(), float(vh or 0.0)))

    total = sum(v for _, v in cuentas)
    if abs(total) > 1.0:
        raise EECCError(
            f"El sumas y saldos no cuadra: saldos deudores menos acreedores = {fmt_ar(total)}. "
            "Revisar la exportación del sistema.")
    if hist and abs(sum(v for _, v in hist)) > 1.0:
        hist = []
    return cuentas, dict(hist=hist or None, ajustado=(c_hist is not None) or solo_ajustado,
                         solo_ajustado=solo_ajustado)


# ════════════════════════════════════════════════════════════════════════════
# Lectura del EECC anterior (PDF)
# ════════════════════════════════════════════════════════════════════════════
def texto_pdf(path):
    """pdftotext -layout o pypdf (modo layout): respetan columnas sin partir los números."""
    import shutil
    import subprocess
    if shutil.which("pdftotext"):
        out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
        if out.strip():
            return out
    try:
        import pypdf
        out = "\n".join(pg.extract_text(extraction_mode="layout") or "" for pg in pypdf.PdfReader(path).pages)
        if out.strip():
            return out
    except Exception:
        pass
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


class Prior:
    """Valores del EECC anterior (columna 'Actual' de ese balance, moneda de su cierre)."""

    LABELS = {
        "caja": r"caja y bancos",
        "inv": r"^\s*inversiones|\binversiones\s*(\(|$|-|\")",
        "cv": r"creditos por ventas",
        "oc": r"otros creditos",
        "bc": r"bienes de cambio",
        "bu": r"bienes de uso",
        "ta": r"total del activo(?! (no )?corriente)",
        "dc": r"comerciales",
        "prest": r"prestamos",
        "rem": r"remuneraciones y cargas sociales",
        "df": r"cargas fiscales",
        "ds": r"deudas sociales|sociales\s*\(",
        "od": r"otras deudas",
        "tp": r"total del pasivo(?! (no )?corriente)(?! y)",
        "pn": r"patrimonio neto",
        "ventas": r"ventas netas",
        "costo": r"costo de (los )?(bienes|mercaderias|servicios)",
        "gcom": r"gastos de comercializacion",
        "gadm": r"gastos de administracion",
        "og": r"^\s*otros gastos",
        "rfyt": r"resultados financieros y por tenencia",
        "oi": r"otros ingresos y egresos",
        "ig": r"^\s*impuesto a las ganancias",
        "res": r"(ganancia|resultado).{0,12}del ejercicio",
        "efe_ini": r"efectivo (y sus equivalentes )?al inicio",
        "efe_cie": r"efectivo (y sus equivalentes )?al cierre",
        "efe_op": r"flujo neto.*actividades operativas",
        "efe_inv": r"flujo neto.*actividades de inversion",
        "efe_fin": r"flujo neto.*actividades de financiacion",
        "efe_dep": r"depreciacion|amortizacion de bienes",
        "efe_cv": r"disminucion\)? en creditos por ventas",
        "efe_oc": r"disminucion\)? en otros creditos",
        "efe_bc": r"disminucion\)? en bienes de cambio",
        "efe_deu": r"disminucion\)? en deudas",
    }

    def __init__(self, path):
        self.text = texto_pdf(path) if path else ""
        self.lines = self.text.splitlines()
        self.nlines = [norm(l) for l in self.lines]
        self.v = {}
        self.warn = []
        if not self.text.strip():
            return
        for k, pat in self.LABELS.items():
            self.v[k] = self._find(pat)
        # notas: cuentas
        self.notas_ini = 0
        for i, nl in enumerate(self.nlines):
            if "notas a los estados contables" in nl and i > 20:
                self.notas_ini = i
                break
        self.anexo_bu = self._leer_anexo_bu()
        self.anexo_gastos = self._leer_anexo_gastos()
        self.normas_rt54 = bool(re.search(r"tecnica\s*(n.{0,3})?\s*54|rt\s*54", norm(self.text)))
        m = re.search(r"actividad principal\s*:?\s*(.+)", self.text, re.I)
        self.actividad = ""
        if m:
            act = m.group(1).strip()
            idx = self.text[:m.end()].count("\n")
            if idx + 1 < len(self.lines):
                nxt = self.lines[idx + 1].strip()
                if nxt and nxt.isupper() and ":" not in nxt and len(nxt) < 60:
                    act += " " + nxt
            self.actividad = re.sub(r"\s+", " ", act)

    def _find(self, pat, start=0):
        rx = re.compile(pat)
        for i in range(start, len(self.lines)):
            m = rx.search(self.nlines[i])
            if not m:
                continue
            toks = tokens_after(self.lines[i], m.end())
            if toks:
                return toks[0]
        return None

    def g(self, k):
        return self.v.get(k) or 0.0

    def cuenta(self, nombre):
        """Saldo de una cuenta en las notas del EECC anterior (match exacto del nombre)."""
        key = skey(nombre)
        for i in range(self.notas_ini, len(self.lines)):
            line = self.lines[i]
            m = NUM_RE.search(line)
            while m and parse_num(m.group()) is None:
                m = NUM_RE.search(line, m.end())
            if not m:
                continue
            label = skey(line[:m.start()])
            if label == key:
                return parse_num(m.group())
        return None

    def _leer_anexo_gastos(self):
        """Renglones del anexo de gastos (art. 64 LGS) del EECC anterior: [(concepto, total)]."""
        ini = None
        for i, nl in enumerate(self.nlines):
            if re.search(r"informacion\s+requerida|art\.?\s*64", nl):
                ini = i
                break
        out = []
        if ini is None:
            return out
        for i in range(ini + 1, min(ini + 60, len(self.lines))):
            nl = self.nlines[i]
            if re.search(r"^\s*totales?\b", nl):
                break
            if re.search(r"notas a los estados|anexo i{1,3}\b", nl) and i > ini + 3:
                break
            m = re.match(r"\s*([a-z][^0-9()]*?)\s{2,}", nl)
            if not m:
                continue
            lab = self.lines[i][m.start(1):m.end(1)].strip()
            if re.search(r"^(detalle|ejercicio|gastos\s*$|total\s*$|admin|comerc|otros\s*gastos\s*$)", nkey(lab)):
                continue
            toks = tokens_after(self.lines[i], m.end())
            if toks and abs(toks[0]) > 0.005:
                out.append((lab, abs(toks[0])))
        return out

    def _leer_anexo_bu(self):
        ini = None
        for i, nl in enumerate(self.nlines):
            if re.search(r"valores\s+de\s+(incorporacion|origen)", nl):
                ini = i
                break
        out = {}
        if ini is None:
            return out
        for i in range(ini, min(ini + 40, len(self.lines))):
            nl = self.nlines[i]
            if re.search(r"anexo ii|costo de (las )?mercaderias|informacion requerida|^\s*totales", nl):
                break
            label_m = re.match(r"\s*([a-z][a-z .y]+?)\s{2,}", nl)
            if not label_m:
                continue
            toks = tokens_after(self.lines[i], label_m.end())
            if len(toks) < 9:
                continue
            tipo = tipo_bu(label_m.group(1))
            origen = toks[3]
            acum = toks[-2]
            neto = toks[-1]
            tasa = toks[6] if len(toks) >= 10 else 0.0
            if 0 < tasa <= 1:
                tasa *= 100
            if origen or neto:
                out[tipo] = dict(origen=origen, acum=acum, neto=neto, tasa=tasa)
        return out


# ════════════════════════════════════════════════════════════════════════════
# Cálculo
# ════════════════════════════════════════════════════════════════════════════
def calcular(a, cuentas, prior):
    COF = a.cof
    rx = lambda v: round((v or 0.0) * COF, 2)
    warn = []
    cierre = datetime.strptime(a.fecha_cierre, "%Y-%m-%d").date()
    try:
        inicio = cierre.replace(year=cierre.year - 1) + timedelta(days=1)
    except ValueError:
        inicio = cierre - timedelta(days=364)
    if a.fecha_inicio:
        inicio = datetime.strptime(a.fecha_inicio, "%Y-%m-%d").date()
    cierre_ant = inicio - timedelta(days=1)
    tiene_prior = bool(prior.text.strip()) and a.nro_ejercicio > 1

    # ── clasificación ──
    clas, desconocidas = [], []
    for nom, sal in cuentas:
        k = clasificar(nom)
        if k == "caja" and sal < -0.5:
            k = "prest"
        if k is None:
            if abs(sal) > 0.005:
                desconocidas.append((nom, sal))
            continue
        clas.append((nom, sal, k))
    if desconocidas:
        lista = "; ".join(f"{n} ({fmt_ar(s)})" for n, s in desconocidas)
        raise EECCError("Cuentas del sumas y saldos que no pude clasificar en un rubro: " + lista)

    def suma(k, signo=1):
        return signo * sum(s for _, s, kk in clas if kk == k)

    moneda_ext = any(MONEDA_EXT.search(nkey(n)) and abs(s) > 0.005 for n, s, _ in clas)

    # ── prior (homogéneo al cierre actual) ──
    P = {k: rx(prior.g(k)) if tiene_prior else 0.0 for k in Prior.LABELS}
    for k in ("costo", "gcom", "gadm", "og", "ig"):
        P[k] = abs(P[k])

    # ── ¿El sistema ya registró el ajuste por inflación? (cuenta RECPAM con saldo, p. ej. asiento de Xubio) ──
    # En ese caso se respetan los saldos del sumas y saldos (PN y bienes de uso ya reexpresados) y no se proponen
    # asientos de reexpresión; los papeles de trabajo quedan como control.
    recpam_ss = round(sum(s for n, s, k in clas if k == "rfyt" and re.search(r"recpam", nkey(n))), 2)
    axi_sistema = abs(recpam_ss) > 1 or bool(getattr(a, "ss_ajustado", False))

    # ── Bienes de uso ──
    bu_ss, bu_orig, bu_acum = {}, {}, {}
    for n, s, k in clas:
        if k == "bu":
            bu_ss.setdefault(tipo_bu(n), 0.0)
            bu_ss[tipo_bu(n)] += s
            dst = bu_acum if re.search(r"amort|depreciacion", nkey(n)) else bu_orig
            dst[tipo_bu(n)] = dst.get(tipo_bu(n), 0.0) + s
    anexo_p = prior.anexo_bu if tiene_prior else {}
    if tiene_prior and P["bu"] and not anexo_p:
        warn.append("No pude leer el Anexo I del EECC anterior: se tomó el neto anterior como valor de origen "
                    "y tasas estándar. Revisar el Anexo de bienes de uso.")
        tot_ss = sum(v for v in bu_ss.values() if v > 0) or 1.0
        tipos = [t for t, v in bu_ss.items() if v > 0] or ["Otros bienes de uso"]
        anexo_p = {t: dict(origen=prior.g("bu") * (bu_ss.get(t, tot_ss) / tot_ss), acum=0.0,
                           neto=prior.g("bu") * (bu_ss.get(t, tot_ss) / tot_ss), tasa=0.0) for t in tipos}
    amort_booked = suma("amort")
    bu_rows = []
    for tipo in sorted(set(bu_ss) | set(anexo_p), key=lambda t: [x[0] for x in BU_TIPOS].index(t)
                       if t in [x[0] for x in BU_TIPOS] else 99):
        p = anexo_p.get(tipo, dict(origen=0.0, acum=0.0, neto=0.0, tasa=0.0))
        tasa = p["tasa"] or tasa_default(tipo)
        net_ss = bu_ss.get(tipo, 0.0)
        altas_nom = net_ss - p["neto"]
        if abs(altas_nom) < 1:
            altas_nom = 0.0
        o_ini, a_ini = rx(p["origen"]), rx(p["acum"])
        altas_h = bajas_h = 0.0
        if altas_nom > 0:
            altas_h = round(altas_nom * COF ** 0.5, 2)
        elif altas_nom < 0:
            bajas_h = round(-altas_nom * COF, 2)
            warn.append(f"Bienes de uso '{tipo}': el saldo bajó respecto del ejercicio anterior; "
                        "se expuso como baja. Verificar.")
        o_cie = o_ini + altas_h - bajas_h
        if axi_sistema:
            # valores de origen y amortización acumulada tal como están en el sistema (ya reexpresados)
            o_cie = round(bu_orig.get(tipo, 0.0), 2)
            altas_h, bajas_h = max(o_cie - o_ini, 0.0), max(o_ini - o_cie, 0.0)
            a_ini = round(-bu_acum.get(tipo, 0.0), 2)
        amort = round(min(o_cie * tasa / 100, max(o_cie - a_ini, 0)), 2) if not amort_booked else 0.0
        bu_rows.append(dict(tipo=tipo, o_ini=o_ini, altas=altas_h, bajas=bajas_h, o_cie=o_cie,
                            a_ini=a_ini, tasa=tasa, amort=amort, a_cie=a_ini + amort,
                            neto=o_cie - a_ini - amort, neto_ant=rx(p["neto"]),
                            o_ant=p["origen"], a_ant=p["acum"], altas_nom=max(altas_nom, 0.0)))
    if amort_booked:
        tot_o = sum(r["o_cie"] for r in bu_rows) or 1
        for r in bu_rows:
            r["amort"] = round(amort_booked * r["o_cie"] / tot_o, 2)
            if axi_sistema:      # la acumulada del sistema ya incluye la amortización registrada
                r["a_ini"] = round(r["a_ini"] - r["amort"], 2)
            r["a_cie"] = r["a_ini"] + r["amort"]
            r["neto"] = r["o_cie"] - r["a_cie"]
    amort_ej = sum(r["amort"] for r in bu_rows)
    altas_tot = sum(r["altas"] for r in bu_rows)
    bajas_tot = sum(r["bajas"] for r in bu_rows)

    # ── Activo / Pasivo al cierre ──
    C = dict(caja=suma("caja"), inv=suma("inv"), cv=suma("cv"), oc=suma("oc"), bc=suma("bc"),
             bu=sum(r["neto"] for r in bu_rows),
             dc=suma("dc", -1), prest=suma("prest", -1), rem=suma("rem", -1), df=suma("df", -1),
             ds=suma("ds", -1), od=suma("od", -1))
    for k in ("cv", "oc", "bc", "inv"):
        if C[k] < -0.5:
            warn.append(f"El rubro {RUBROS[k][1]} quedó con saldo negativo ({fmt_ar(C[k])}).")
    C["ac"] = C["caja"] + C["inv"] + C["cv"] + C["oc"] + C["bc"]
    C["anc"] = C["bu"]
    C["ta"] = C["ac"] + C["anc"]
    C["tp"] = C["dc"] + C["prest"] + C["rem"] + C["df"] + C["ds"] + C["od"]
    C["pn"] = C["ta"] - C["tp"]

    # ── Patrimonio neto de inicio ──
    cap_ss = suma("cap", -1)
    cap = a.cap_nominal if a.cap_nominal else cap_ss
    if a.cap_nominal and abs(cap_ss - a.cap_nominal) > 1:
        warn.append(f"El capital del sumas y saldos ({fmt_ar(cap_ss)}) no coincide con el cargado en Notion "
                    f"({fmt_ar(a.cap_nominal)}). Se usó el de Notion.")
    aj_p, apirr_p, res_p, rna_p = suma("aj", -1), suma("apirr", -1), suma("resleg", -1), suma("rna", -1)
    pn_ss = cap_ss + aj_p + apirr_p + res_p + rna_p
    ss_pn = dict(cap=cap_ss, aj=aj_p, apirr=apirr_p, resleg=res_p, rna=rna_p)
    aportes_nom = 0.0
    aj_base = aj_p
    pn_control = None
    if tiene_prior and axi_sistema:
        # PN de inicio = el que dejó el sistema (ya reexpresado); se controla contra el PN anterior × coeficiente
        pn_prior = prior.g("pn")
        pn_control = round(pn_prior * COF, 2)
        aj_base = aj_p - (pn_ss - pn_prior)
        I = dict(cap=cap_ss, aj=aj_p, apirr=apirr_p, resleg=res_p, rna=rna_p)
        pn_base = dict(cap=cap_ss, aj=aj_base, apirr=apirr_p, resleg=res_p,
                       rna=pn_prior - (cap_ss + aj_base + apirr_p + res_p))
        warn.append(f"El sistema ya tiene registrado el ajuste por inflación (RECPAM {fmt_ar(recpam_ss)}): se "
                    "respetaron sus saldos y no se proponen asientos de reexpresión.")
        if abs(pn_ss - pn_control) > max(1.0, abs(pn_control) * 0.005):
            warn.append(f"El PN de inicio reexpresado por el sistema ({fmt_ar(pn_ss)}) difiere del PN anterior × "
                        f"coeficiente de Notion ({fmt_ar(pn_control)}). Revisar el coeficiente usado en Xubio.")
    elif tiene_prior:
        pn_prior = prior.g("pn")
        dif = pn_ss - pn_prior
        if abs(dif) > max(1.0, abs(pn_prior) * 0.001):
            if dif > 0 and abs(cap_ss - cap) < 1:
                aportes_nom = dif
                warn.append(f"El PN del sumas y saldos supera al del EECC anterior en {fmt_ar(dif)}; "
                            "se expuso como aporte del ejercicio. Verificar.")
            else:
                warn.append(f"El PN del sumas y saldos ({fmt_ar(pn_ss)}) no coincide con el PN del EECC "
                            f"anterior ({fmt_ar(pn_prior)}). Diferencia {fmt_ar(dif)}.")
        # inicio = PN anterior (sin los aportes nuevos)
        base_cap, base_aj, base_apirr, base_res = cap, aj_p, apirr_p, res_p
        base_rna = pn_prior - (cap + aj_p + apirr_p + res_p) if aportes_nom == 0 else rna_p
        I = dict(cap=cap, aj=round((base_cap + base_aj) * COF - base_cap, 2), apirr=rx(base_apirr),
                 resleg=rx(base_res), rna=rx(base_rna))
        pn_base = dict(cap=base_cap, aj=base_aj, apirr=base_apirr, resleg=base_res, rna=base_rna)
    else:
        pn_base = dict(cap=0.0, aj=0.0, apirr=0.0, resleg=0.0, rna=0.0)
        aportes_nom = pn_ss - aj_p      # si el sistema ya registró ajuste de capital, se recalcula sobre el nominal
        I = dict(cap=0.0, aj=0.0, apirr=0.0, resleg=0.0, rna=0.0)
    I["aportes"] = I["cap"] + I["aj"] + I["apirr"]
    I["pn"] = I["aportes"] + I["resleg"] + I["rna"]
    # Primer ejercicio: el capital se suscribe en la constitución (inicio del ejercicio) y se reexpresa con el
    # coeficiente del ejercicio; el ajuste va a «Ajuste de capital».
    aportes_aj = 0.0
    if not tiene_prior:
        aportes_aj = round(aj_p, 2) if axi_sistema else round(aportes_nom * (COF - 1), 2)
    aportes_h = round(aportes_nom + aportes_aj, 2)
    integ_pend = sum(s for n, s, k in clas if k == "oc" and re.search(r"integra|suscripcion pendiente", nkey(n)))

    # ── Resultados ──
    ventas = suma("ventas", -1)
    compras = suma("costo")
    bc_ini, bc_cie = P["bc"], C["bc"]
    usa_inventario = bool(bc_ini or bc_cie)
    costo = (bc_ini + compras - bc_cie) if usa_inventario else compras
    gcom = suma("gcom")
    gadm_ss = suma("gadm")
    gadm = gadm_ss + amort_ej + (amort_booked if amort_booked else 0)
    oi = suma("oi", -1)
    ig = suma("ig")
    res = C["pn"] - I["pn"] - aportes_h
    rfyt = round(res - (ventas - costo - gcom - gadm + oi - ig), 2)
    R = dict(ventas=ventas, costo=costo, bruta=ventas - costo, gcom=gcom, gadm=gadm, og=0.0,
             rfyt=rfyt, oi=oi, antes_ig=res + ig, ig=ig, res=res, compras=compras,
             bc_ini=bc_ini, bc_cie=bc_cie, usa_inv=usa_inventario)

    # ── Comparativos (ejercicio anterior reexpresado) ──
    A = dict(P)
    if tiene_prior:
        A["ac"] = A["caja"] + A["inv"] + A["cv"] + A["oc"] + A["bc"]
        A["anc"] = A["bu"]
        A["ta"] = A["ta"] or (A["ac"] + A["anc"])
        A["tp"] = A["dc"] + A["prest"] + A["rem"] + A["df"] + A["ds"] + A["od"]
        A["pn"] = A["pn"] or (A["ta"] - A["tp"])
        A["bruta"] = A["ventas"] - A["costo"]
        A["antes_ig"] = A["res"] + A["ig"]
        if abs((A["ta"] - A["tp"]) - A["pn"]) > max(2.0, abs(A["pn"]) * 0.001):
            warn.append("En el EECC anterior no pude reconstruir Activo − Pasivo = PN; revisar comparativos.")

    # ── EFE comparativo ──
    AE = {}
    if tiene_prior:
        if a.nro_ejercicio - 1 == 1:
            # el ejercicio anterior fue el primero: saldos de inicio en cero, se reconstruye completo
            dep = sum(rx(p["acum"]) for p in anexo_p.values())
            AE = dict(ini=0.0, cie=A["caja"], res=A["res"], amort=dep, cv=-A["cv"], oc=-A["oc"], bc=-A["bc"],
                      deu=A["dc"] + A["rem"] + A["df"] + A["od"], bu=-(A["bu"] + dep), invs=-A["inv"],
                      aportes=A["pn"] - A["res"], prest=A["prest"], ds=A["ds"])
            AE["op"] = AE["res"] + AE["amort"] + AE["cv"] + AE["oc"] + AE["bc"] + AE["deu"]
            AE["inv"] = AE["bu"] + AE["invs"]
            AE["fin"] = AE["aportes"] + AE["prest"] + AE["ds"]
        else:
            m = dict(ini="efe_ini", cie="efe_cie", res="res", amort="efe_dep", cv="efe_cv", oc="efe_oc",
                     bc="efe_bc", deu="efe_deu", op="efe_op", inv="efe_inv", fin="efe_fin")
            AE = {k: (A.get(v) if prior.v.get(v) is not None else None) for k, v in m.items()}
            AE["res"] = A["res"]
        AE["var"] = (AE.get("cie") or 0) - (AE.get("ini") or 0)
        if AE.get("op") is not None and AE.get("inv") is not None and AE.get("fin") is not None:
            if abs(AE["op"] + AE["inv"] + AE["fin"] - AE["var"]) > max(2.0, abs(AE["var"]) * 0.01):
                warn.append("El flujo de efectivo del EECC anterior no cierra; revisar la columna comparativa.")

    # ── EFE (método indirecto, moneda homogénea) ──
    d = lambda k: C[k] - A.get(k, 0.0)
    E = dict(ini=A.get("caja", 0.0), cie=C["caja"])
    E["var"] = E["cie"] - E["ini"]
    E["res"] = res
    E["amort"] = amort_ej
    E["cv"] = -d("cv")
    E["oc"] = -(d("oc") - (integ_pend if not tiene_prior else 0.0))
    E["bc"] = -d("bc")
    E["deu"] = d("dc") + d("rem") + d("df") + d("od")
    E["op"] = E["res"] + E["amort"] + E["cv"] + E["oc"] + E["bc"] + E["deu"]
    E["bu"] = -(C["bu"] + amort_ej - A.get("bu", 0.0))
    E["invs"] = -d("inv")
    E["inv"] = E["bu"] + E["invs"]
    E["aportes"] = aportes_h - (integ_pend if not tiene_prior else 0.0)
    E["prest"] = d("prest")
    E["ds"] = d("ds")
    E["fin"] = E["aportes"] + E["prest"] + E["ds"]
    E["chk"] = E["op"] + E["inv"] + E["fin"] - E["var"]
    if abs(E["chk"]) > 1:
        raise EECCError(f"El flujo de efectivo no cierra por {fmt_ar(E['chk'])} (error interno).")

    # ── Controles ──
    if abs(C["ta"] - C["tp"] - C["pn"]) > 0.01:
        raise EECCError("Activo − Pasivo ≠ PN (error interno).")
    if abs(I["pn"] + aportes_h + res - C["pn"]) > 0.01:
        raise EECCError("La evolución del PN no cierra (error interno).")
    if C["pn"] < 0:
        warn.append(f"El patrimonio neto al cierre es negativo ({fmt_ar(C['pn'])}).")

    return dict(cuentas=clas, C=C, A=A, I=I, R=R, E=E, AE=AE, bu=bu_rows, amort_ej=amort_ej, pn_base=pn_base, ss_pn=ss_pn, amort_booked=amort_booked,
                axi_sistema=axi_sistema, recpam_ss=recpam_ss, pn_control=pn_control,
                aportes=aportes_h, aportes_nom=round(aportes_nom, 2), aportes_aj=aportes_aj, integ_pend=integ_pend,
                inicio=inicio, cierre=cierre, cierre_ant=cierre_ant,
                tiene_prior=tiene_prior, moneda_ext=moneda_ext, warn=warn, cof=COF,
                primer_rt54=not (tiene_prior and prior.normas_rt54),
                actividad=a.actividad or prior.__dict__.get("actividad", ""),
                prior=prior)


# ════════════════════════════════════════════════════════════════════════════
# Armado de los estados (modelo de filas común a XLSX y PDF)
# ════════════════════════════════════════════════════════════════════════════
def armar(a, X):
    C, A, I, R, E = X["C"], X["A"], X["I"], X["R"], X["E"]
    comp = X["tiene_prior"]
    hdr2 = ["Actual", "Anterior"] if comp else ["Actual"]
    v = lambda c, k: [C.get(k, 0.0), A.get(k, 0.0)] if comp else [C.get(k, 0.0)]

    # numeración de anexos
    anexos = []
    if X["bu"]:
        anexos.append("bu")
    if R["usa_inv"]:
        anexos.append("costo")
    anexos.append("gastos")
    ROM = ["I", "II", "III", "IV"]
    anx = {k: ROM[i] for i, k in enumerate(anexos)}

    # numeración de notas de composición (3.x)
    comp_rubros = []
    for k in ["caja", "inv", "cv", "oc", "bc", "dc", "prest", "rem", "df", "ds", "od"]:
        if abs(C.get(k, 0)) > 0.005 or abs(A.get(k, 0)) > 0.005:
            comp_rubros.append(k)
    comp_rubros.append("rfyt")
    nota_n = {k: f"3.{i + 1}" for i, k in enumerate(comp_rubros)}
    nref = lambda k: f"(Nota {nota_n[k]})" if k in nota_n else ""

    S = []

    # ── ESP ──
    rows = []
    rows.append(("sec", "ACTIVO", None))
    rows.append(("sub", "Activo corriente", None))
    for k, lab in [("caja", "Caja y bancos"), ("inv", "Inversiones"), ("cv", "Créditos por ventas"),
                   ("oc", "Otros créditos"), ("bc", "Bienes de cambio")]:
        rows.append(("det", f"{lab} {nref(k)}".strip(), v(C, k)))
    rows.append(("tot", "Total del activo corriente", v(C, "ac")))
    rows.append(("sub", "Activo no corriente", None))
    rows.append(("det", f"Bienes de uso (Anexo {anx['bu']})" if "bu" in anx else "Bienes de uso", v(C, "bu")))
    rows.append(("tot", "Total del activo no corriente", v(C, "anc")))
    rows.append(("grand", "TOTAL DEL ACTIVO", v(C, "ta")))
    rows.append(("blank", "", None))
    rows.append(("sec", "PASIVO", None))
    rows.append(("sub", "Pasivo corriente", None))
    for k, lab in [("dc", "Deudas comerciales"), ("prest", "Préstamos"),
                   ("rem", "Remuneraciones y cargas sociales"), ("df", "Cargas fiscales"),
                   ("ds", "Deudas con socios"), ("od", "Otras deudas")]:
        rows.append(("det", f"{lab} {nref(k)}".strip(), v(C, k)))
    rows.append(("tot", "Total del pasivo corriente", v(C, "tp")))
    rows.append(("grand", "TOTAL DEL PASIVO", v(C, "tp")))
    rows.append(("blank", "", None))
    rows.append(("grand", "PATRIMONIO NETO (según el EEPN)", v(C, "pn")))
    rows.append(("grand", "TOTAL DEL PASIVO Y PATRIMONIO NETO", [x + y for x, y in zip(v(C, "tp"), v(C, "pn"))]))
    S.append(dict(key="ESP", title="ESTADO DE SITUACIÓN PATRIMONIAL", cols=hdr2, rows=rows, land=False))

    # ── ER ──
    rr = lambda k: [R[k], A.get(k, 0.0)] if comp else [R[k]]
    neg = lambda l: [-x for x in l]
    rows = [
        ("det", "Ventas netas de bienes (o servicios)", rr("ventas")),
        ("det", "Costo de los bienes vendidos (o servicios prestados)" +
         (f" (Anexo {anx['costo']})" if "costo" in anx else ""), neg(rr("costo"))),
        ("tot", "Ganancia (Pérdida) bruta", rr("bruta")),
        ("det", f"Gastos de comercialización (Anexo {anx['gastos']})", neg(rr("gcom"))),
        ("det", f"Gastos de administración (Anexo {anx['gastos']})", neg(rr("gadm"))),
        ("det", "Otros ingresos y egresos", rr("oi")),
        ("det", f"Resultados financieros y por tenencia (incluido el RECPAM) {nref('rfyt')}", rr("rfyt")),
        ("tot", "Ganancia (Pérdida) antes del impuesto a las ganancias", rr("antes_ig")),
        ("det", "Impuesto a las ganancias", neg(rr("ig"))),
        ("grand", "GANANCIA (PÉRDIDA) DEL EJERCICIO", rr("res")),
    ]
    S.append(dict(key="ER", title="ESTADO DE RESULTADOS", cols=hdr2, rows=rows, land=False))

    # ── EEPN ──
    cols = ["Capital social", "Ajuste de capital", "Aportes irrevocables", "Total aportes",
            "Reservas", "Resultados no asignados", "Total actual"] + (["Total anterior"] if comp else [])
    # comparativo (ejercicio anterior)
    if comp:
        prior_first = a.nro_ejercicio - 1 == 1
        a_res = A.get("res", 0.0)
        a_pn = A.get("pn", 0.0)
        a_ini = 0.0 if prior_first else a_pn - a_res
        a_apo = (a_pn - a_res) if prior_first else 0.0

    def eep(det, cap, aj, apirr, resl, rna, ant=None):
        tot_apo = cap + aj + apirr
        vals = [cap, aj, apirr, tot_apo, resl, rna, tot_apo + resl + rna]
        if comp:
            vals.append(ant)
        return vals

    rows = []
    rows.append(("det", "Saldos al inicio del ejercicio",
                 eep("", I["cap"], I["aj"], I["apirr"], I["resleg"], I["rna"], a_ini if comp else None)))
    rows.append(("det", "Modificación de saldos de inicio", eep("", 0, 0, 0, 0, 0, 0.0 if comp else None)))
    rows.append(("tot", "Saldos al inicio modificados",
                 eep("", I["cap"], I["aj"], I["apirr"], I["resleg"], I["rna"], a_ini if comp else None)))
    if X["aportes"] or (comp and a_apo):
        rows.append(("det", "Suscripción de capital / aportes",
                     eep("", X["aportes"] - X["aportes_aj"], X["aportes_aj"], 0, 0, 0, a_apo if comp else None)))
    rows.append(("det", "Ganancia (Pérdida) del ejercicio",
                 eep("", 0, 0, 0, 0, R["res"], A.get("res", 0.0) if comp else None)))
    fin = eep("", I["cap"] + X["aportes"] - X["aportes_aj"], I["aj"] + X["aportes_aj"], I["apirr"], I["resleg"], I["rna"] + R["res"],
              A.get("pn", 0.0) if comp else None)
    rows.append(("grand", "Saldos al cierre del ejercicio", fin))
    S.append(dict(key="EEPN", title="ESTADO DE EVOLUCIÓN DEL PATRIMONIO NETO", cols=cols, rows=rows,
                  land=True))

    # ── EFE ──
    AE = X["AE"]
    ee = lambda k, pk=None: [E[k], AE.get(k)] if comp else [E[k]]
    rows = [
        ("sec", "VARIACIÓN NETA DEL EFECTIVO Y SUS EQUIVALENTES", None),
        ("det", "Efectivo y sus equivalentes al inicio del ejercicio", ee("ini", "efe_ini")),
        ("det", "Efectivo y sus equivalentes al cierre del ejercicio", ee("cie", "efe_cie")),
        ("tot", "Aumento (Disminución) neto del efectivo y sus equivalentes",
         ee("var")),
        ("blank", "", None),
        ("sec", "CAUSAS DE LAS VARIACIONES DEL EFECTIVO Y SUS EQUIVALENTES", None),
        ("sub", "Actividades operativas", None),
        ("det", "Ganancia (Pérdida) del ejercicio", ee("res", "res")),
        ("det", "Ajustes para arribar al flujo neto de efectivo de las actividades operativas:", None),
        ("det", "   Depreciación de bienes de uso", ee("amort", "efe_dep")),
        ("det", "Cambios en activos y pasivos operativos:", None),
        ("det", "   (Aumento) Disminución en créditos por ventas", ee("cv", "efe_cv")),
        ("det", "   (Aumento) Disminución en otros créditos", ee("oc", "efe_oc")),
        ("det", "   (Aumento) Disminución en bienes de cambio", ee("bc", "efe_bc")),
        ("det", "   Aumento (Disminución) en deudas operativas", ee("deu", "efe_deu")),
        ("tot", "Flujo neto de efectivo generado por (utilizado en) las actividades operativas", ee("op", "efe_op")),
        ("sub", "Actividades de inversión", None),
        ("det", "Pagos por compras de bienes de uso", ee("bu")),
        ("det", "(Aumento) Disminución de inversiones no equivalentes de efectivo", ee("invs")),
        ("tot", "Flujo neto de efectivo generado por (utilizado en) las actividades de inversión", ee("inv", "efe_inv")),
        ("sub", "Actividades de financiación", None),
        ("det", "Aportes de capital", ee("aportes")),
        ("det", "Aumento (Disminución) de préstamos", ee("prest")),
        ("det", "Aumento (Disminución) de deudas con socios", ee("ds")),
        ("tot", "Flujo neto de efectivo generado por (utilizado en) las actividades de financiación", ee("fin", "efe_fin")),
        ("grand", "Aumento (Disminución) neto del efectivo y sus equivalentes", ee("var")),
    ]
    S.append(dict(key="EFE", title="ESTADO DE FLUJO DE EFECTIVO (método indirecto)", cols=hdr2, rows=rows,
                  land=False))

    # ── Anexo Bienes de uso ──
    if "bu" in anx:
        cols = ["Al inicio", "Altas", "Bajas", "Al cierre", "Acum. al inicio", "Tasa %",
                "Del ejercicio", "Acum. al cierre", "Neto actual"] + (["Neto anterior"] if comp else [])
        rows = []
        tot = [0.0] * len(cols)
        for r in X["bu"]:
            vals = [r["o_ini"], r["altas"], r["bajas"], r["o_cie"], r["a_ini"], None, r["amort"], r["a_cie"],
                    r["neto"]] + ([r["neto_ant"]] if comp else [])
            rows.append(("det", r["tipo"], vals, {5: f"{r['tasa']:g}".replace(".", ",")}))
            tot = [t + (x or 0) for t, x in zip(tot, vals)]
        tot[5] = None
        rows.append(("grand", "TOTALES", tot))
        S.append(dict(key=f"Anexo {anx['bu']}", title=f"ANEXO {anx['bu']} — BIENES DE USO", cols=cols,
                      rows=rows, land=True))

    # ── Anexo Costo ──
    if "costo" in anx:
        rows = [("det", "Existencia al inicio del ejercicio", [R["bc_ini"]]),
                ("det", "Compras del ejercicio", [R["compras"]]),
                ("det", "Existencia al cierre del ejercicio", [-R["bc_cie"]]),
                ("grand", "Costo de los bienes vendidos", [R["costo"]])]
        S.append(dict(key=f"Anexo {anx['costo']}",
                      title=f"ANEXO {anx['costo']} — COSTO DE LOS BIENES VENDIDOS", cols=["Actual"],
                      rows=rows, land=False))

    # ── Anexo Gastos (art. 64 inc. b) LGS) ──
    rows = []
    gastos = {}
    for n, s_, k in X["cuentas"]:
        if k in ("gcom", "gadm") and abs(s_) > 0.005:
            gastos.setdefault(n, [0.0, 0.0])
            gastos[n][0 if k == "gadm" else 1] += s_
    if X["amort_ej"]:
        gastos["Amortización bienes de uso"] = [X["amort_ej"], 0.0]
    if not R["usa_inv"] and R["costo"]:
        costo_rows = [(n, s_) for n, s_, k in X["cuentas"] if k == "costo" and abs(s_) > 0.005]
    else:
        costo_rows = []

    # comparativo: renglones del anexo anterior, reexpresados y apareados por nombre
    prev = {}
    if comp:
        for lab, v in X["prior"].anexo_gastos:
            kk = skey(lab)
            if re.search(r"amortizacion", kk):
                kk = skey("Amortización bienes de uso")
            prev[kk] = (lab, prev.get(kk, ("", 0.0))[1] + round(v * X["cof"], 2))
        tot_prev_ok = A.get("costo", 0) + A.get("gcom", 0) + A.get("gadm", 0) + A.get("og", 0)
        if prev and abs(sum(v for _, v in prev.values()) - tot_prev_ok) > max(2.0, tot_prev_ok * 0.005):
            X["warn"].append("El anexo de gastos del EECC anterior no suma igual que su estado de resultados; "
                             "revisar la columna comparativa del anexo.")
    usados = set()

    def ant(n):
        kk = skey(n)
        if kk in prev:
            usados.add(kk)
            return prev[kk][1]
        return 0.0

    ncol = 5 if comp else 4
    ttot = [0.0] * ncol
    for n, s_ in costo_rows:
        vals = [s_, 0.0, 0.0, s_] + ([ant(n)] if comp else [])
        rows.append(("det", n, vals))
        ttot = [t + x for t, x in zip(ttot, vals)]
    for n, (ga, gc) in gastos.items():
        vals = [0.0, ga, gc, ga + gc] + ([ant(n)] if comp else [])
        rows.append(("det", n, vals))
        ttot = [t + x for t, x in zip(ttot, vals)]
    if comp:
        for kk, (lab, v) in prev.items():
            if kk not in usados:
                rows.append(("det", lab, [0.0, 0.0, 0.0, 0.0, v]))
                ttot[4] += v
        if not prev:
            ttot[4] = A.get("costo", 0) + A.get("gcom", 0) + A.get("gadm", 0) + A.get("og", 0)
    rows.append(("grand", "TOTALES", ttot))
    S.append(dict(key=f"Anexo {anx['gastos']}",
                  title=f"ANEXO {anx['gastos']} — INFORMACIÓN REQUERIDA POR EL ART. 64 INC. b) LEY 19.550",
                  cols=["Costo de servicios", "Gastos de administración", "Gastos de comercialización",
                        "Total actual"] + (["Total anterior"] if comp else []), rows=rows, land=True))

    notas = armar_notas(a, X, comp_rubros, nota_n)
    return S, notas, anx, nota_n


def armar_notas(a, X, comp_rubros, nota_n):
    C, A, R = X["C"], X["A"], X["R"]
    comp = X["tiene_prior"]
    emp = a.empresa
    N = []  # bloques: ("h1"|"h2"|"p"|"table", contenido)
    N.append(("h1", "1. Notas generales"))
    N.append(("h2", "1.1. Bases de preparación de los estados contables"))
    N.append(("p", "Los presentes estados contables están expresados en pesos en moneda homogénea de cierre "
                   "(Nota 1.3) y han sido preparados de conformidad con la Resolución Técnica N° 54 - T.O. RT 59 "
                   "y modificatorias, emitida por la Federación Argentina de Consejos Profesionales de Ciencias "
                   "Económicas (FACPCE), adoptada por la Resolución P. N° 460/2024 del Consejo Profesional de "
                   "Ciencias Económicas de la Ciudad Autónoma de Buenos Aires (CPCECABA)."))
    ini = X["inicio"].strftime("%d/%m/%Y")
    if a.nro_ejercicio <= 1:
        N.append(("p", f"Por tratarse del primer ejercicio económico de la entidad, iniciado el {ini}, "
                       "no se presenta información comparativa."))
    elif X["primer_rt54"]:
        N.append(("p", f"Esta resolución fue aplicada por la entidad a partir del {ini}, utilizando el enfoque "
                       "retroactivo simplificado, no surgiendo impactos patrimoniales ni en resultados como "
                       "consecuencia de su aplicación. Los estados contables comparativos se exponen reexpresados "
                       "en moneda homogénea al cierre del ejercicio actual conforme a lo establecido en la RT 54."))
    else:
        N.append(("p", "Los estados contables comparativos se exponen reexpresados en moneda homogénea al "
                       "cierre del ejercicio actual conforme a lo establecido en la RT 54."))
    N.append(("h2", "1.2. Clasificación de la entidad"))
    N.append(("p", "De acuerdo con lo establecido por la referida norma, la entidad reviste el carácter de entidad "
                   "pequeña (EP) y no presenta temas de alta complejidad contable en función de la información "
                   "suministrada por la administración."))
    N.append(("h2", "1.3. Unidad de medida"))
    N.append(("p", "Los presentes estados contables han sido preparados en moneda homogénea de cierre, "
                   "reconociendo los efectos de la pérdida del poder adquisitivo de la moneda de acuerdo con las "
                   "normas contables profesionales vigentes. En la preparación se utilizaron las simplificaciones "
                   "admitidas para entidades pequeñas, exponiendo los resultados financieros y por tenencia, "
                   "incluyendo el RECPAM, en una sola línea del estado de resultados."))
    N.append(("h2", "1.4. Uso de estimaciones en la preparación de los presentes estados contables"))
    N.append(("p", f"La preparación de estados contables requiere que el órgano de administración de {emp} "
                   "realice estimaciones y evaluaciones que afectan el monto de los activos y pasivos registrados "
                   "y los activos y pasivos contingentes revelados a la fecha de cierre, como así también los "
                   "ingresos y egresos registrados en el ejercicio. Los resultados reales futuros pueden diferir "
                   "de las estimaciones y evaluaciones realizadas a la fecha de preparación de los presentes "
                   "estados contables."))

    # 2. Políticas
    N.append(("h1", "2. Políticas contables sobre criterios de medición contable"))
    pol = []
    tiene = lambda k: abs(C.get(k, 0)) > 0.005 or abs(A.get(k, 0)) > 0.005
    if X["moneda_ext"]:
        pol.append(("Activos y pasivos en moneda extranjera",
                    "Los activos y pasivos en moneda extranjera han sido medidos al tipo de cambio "
                    "correspondiente a la fecha de cierre de los estados contables o, en su defecto, al tipo de "
                    "cambio implícito utilizado por la entidad financiera para su liquidación o cancelación."))
    pol.append(("Caja y bancos", "El efectivo disponible y los saldos en bancos han sido medidos por su importe "
                                 "nominal."))
    if tiene("inv"):
        pol.append(("Inversiones", "Las inversiones temporarias han sido medidas a su valor neto de realización "
                                   "a la fecha de cierre."))
    pol.append(("Cuentas por cobrar a clientes en moneda, Créditos impositivos, Créditos en moneda con partes "
                "relacionadas, Otras cuentas por cobrar en moneda",
                "Las cuentas por cobrar a clientes en moneda, créditos impositivos y otras cuentas por cobrar "
                "se miden a su valor nominal, netos de previsiones por desvalorización, de corresponder."))
    if tiene("bc"):
        pol.append(("Bienes de cambio",
                    "Los bienes de cambio adquiridos se miden a su costo de la última compra, excepto por los "
                    "productos en proceso o terminados que se miden al costo de producción a la fecha de cierre "
                    "de cada ejercicio.\nEl valor registrado de los bienes de cambio no supera su valor recuperable "
                    "estimado en las condiciones actuales. Todos los bienes se encuentran reexpresados a la fecha "
                    "de cierre de acuerdo a lo indicado en la Nota 1.3."))
    if X["bu"]:
        pol.append(("Bienes de uso",
                    "Los bienes de uso se miden a su costo de adquisición reexpresado de acuerdo con lo indicado "
                    "en la Nota 1.3, neto de las depreciaciones acumuladas. Las depreciaciones se calculan por el "
                    "método de la línea recta, aplicando tasas anuales suficientes para extinguir los valores al "
                    "final de la vida útil estimada de los bienes, computándose el año de alta completo. El valor "
                    "registrado de los bienes de uso no supera su valor recuperable."))
    pol.append(("Proveedores de bienes y servicios, Deudas fiscales, Deudas laborales y previsionales, y Deudas "
                "en moneda con partes relacionadas",
                "Las deudas con proveedores y otros pasivos financieros se miden por su valor nominal" +
                (", excepto por los pasivos en moneda extranjera, que se miden de acuerdo con el criterio "
                 "indicado en la Nota 2.1." if X["moneda_ext"] else ".")))
    if tiene("bc"):
        pol.append(("Costo de los bienes vendidos",
                    "El costo de los bienes vendidos se calcula aplicando la técnica de diferencia de inventario "
                    "considerando la reexpresión de la Nota 1.3."))
    pol.append(("Impuesto a las ganancias",
                "El cargo contable por impuesto a las ganancias se reconoce en función del impuesto determinado, "
                "considerando la normativa fiscal vigente."))
    pol.append(("Cuentas de Patrimonio neto",
                "Se encuentran reexpresadas de acuerdo con lo mencionado en la Nota 1.3. Las cuentas “Capital” y "
                "“Aportes Irrevocables” se mantienen a sus valores nominales y el ajuste derivado de dicha "
                "reexpresión monetaria realizada de acuerdo a la Nota 1.3, se expone en las cuentas de “Ajuste de "
                "capital” y “Ajuste de Aportes Irrevocables”, respectivamente."))
    for i, (t, p) in enumerate(pol):
        N.append(("h2", f"2.{i + 1}. {t}"))
        for par in p.split("\n"):
            N.append(("p", par))

    # 3. Composición
    N.append(("h1", "3. Composición de los principales rubros"))
    titulos = dict(caja="Caja y bancos", inv="Inversiones", cv="Créditos por ventas", oc="Otros créditos",
                   bc="Bienes de cambio", dc="Deudas comerciales", prest="Préstamos",
                   rem="Remuneraciones y cargas sociales", df="Cargas fiscales", ds="Deudas con socios",
                   od="Otras deudas", rfyt="Resultados financieros y por tenencia")
    prior = X["prior"]
    signo = lambda k: -1 if RUBROS.get(k, ("A",))[0] == "P" else 1
    for k in comp_rubros:
        filas = []
        if k == "rfyt":
            filas.append(("RECPAM y otros resultados financieros y por tenencia",
                          R["rfyt"], A.get("rfyt", 0.0) if comp else None))
            tot = (R["rfyt"], A.get("rfyt", 0.0) if comp else None)
        else:
            sum_prev = 0.0
            for n, s, kk in X["cuentas"]:
                if kk != k:
                    continue
                act = s * signo(k)
                ant = None
                if comp:
                    pv = prior.cuenta(n)
                    ant = round(pv * X["cof"], 2) if pv is not None else 0.0
                    sum_prev += ant
                if abs(act) > 0.005 or (ant and abs(ant) > 0.005):
                    filas.append((n, act, ant))
            if comp:
                dif = A.get(k, 0.0) - sum_prev
                if abs(dif) > 1:
                    filas.append(("Saldos del ejercicio anterior en otras cuentas", 0.0, dif))
            tot = (C.get(k, 0.0), A.get(k, 0.0) if comp else None)
        N.append(("h2", f"{nota_n[k]}. {titulos[k]}"))
        N.append(("table", dict(cols=["Actual", "Anterior"] if comp else ["Actual"],
                                rows=[(f, [x, y] if comp else [x]) for f, x, y in filas],
                                total=[tot[0], tot[1]] if comp else [tot[0]])))

    if C["pn"] < 0:
        N.append(("h1", "4. Patrimonio neto negativo"))
        N.append(("p", f"Al {X['cierre'].strftime('%d/%m/%Y')} la sociedad presenta un patrimonio neto negativo de "
                       f"$ {fmt_ar(C['pn'])}. Los socios han manifestado su intención de continuar con la "
                       "actividad y de adoptar las medidas necesarias para recomponer el patrimonio, por lo que los "
                       "presentes estados contables han sido preparados bajo el principio de empresa en "
                       "funcionamiento."))
    return N


# ════════════════════════════════════════════════════════════════════════════
# Salida XLSX
# ════════════════════════════════════════════════════════════════════════════
def escribir_xlsx(path, a, X, S, notas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    NF = '#,##0.00;(#,##0.00);"-"'
    BOLD = Font(bold=True)
    HDR = PatternFill("solid", fgColor="1F3864")
    TOT = PatternFill("solid", fgColor="D9E1F2")
    GR = PatternFill("solid", fgColor="B4C6E7")
    thin = Side(style="thin", color="999999")

    wb = Workbook()
    ws = wb.active
    ws.title = "Carátula"
    car = caratula_lineas(a, X)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 80
    for i, (k, val) in enumerate(car, 1):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, val)

    for st in S:
        ws = wb.create_sheet(st["key"][:31])
        ncol = len(st["cols"])
        ws.column_dimensions["A"].width = 62
        for j in range(ncol):
            ws.column_dimensions[get_column_letter(j + 2)].width = 17
        ws.cell(1, 1, a.empresa).font = Font(bold=True, size=12)
        ws.cell(2, 1, st["title"]).font = Font(bold=True, size=11)
        ws.cell(3, 1, subtitulo(a, X)).font = Font(italic=True, size=9)
        r = 5
        for j, c in enumerate(st["cols"]):
            cell = ws.cell(r, j + 2, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HDR
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.cell(r, 1, "").fill = HDR
        ws.row_dimensions[r].height = 30
        for row in st["rows"]:
            r += 1
            kind, lab, vals = row[0], row[1], row[2]
            txt = row[3] if len(row) > 3 else {}
            c = ws.cell(r, 1, lab)
            if kind in ("sec", "sub", "tot", "grand"):
                c.font = BOLD
            if kind == "det":
                c.alignment = Alignment(indent=1)
            for j, x in enumerate(vals or []):
                cell = ws.cell(r, j + 2)
                if j in txt:
                    cell.value = txt[j]
                    cell.alignment = Alignment(horizontal="center")
                elif x is not None:
                    cell.value = round(x, 2)
                    cell.number_format = NF
                if kind in ("tot", "grand"):
                    cell.font = BOLD
                    cell.border = Border(top=thin)
            if kind == "tot":
                for j in range(ncol + 1):
                    ws.cell(r, j + 1).fill = TOT
            if kind == "grand":
                for j in range(ncol + 1):
                    ws.cell(r, j + 1).fill = GR
        r += 2
        ws.cell(r, 1, "Las notas y anexos que se acompañan forman parte integrante de los estados contables.") \
            .font = Font(italic=True, size=8)
        ws.page_setup.orientation = "landscape" if st["land"] else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    r = 1
    ws.cell(r, 1, f"{a.empresa} — NOTAS A LOS ESTADOS CONTABLES").font = Font(bold=True, size=12)
    r += 1
    ws.cell(r, 1, subtitulo(a, X)).font = Font(italic=True, size=9)
    for kind, cont in notas:
        r += 1
        if kind == "h1":
            r += 1
            ws.cell(r, 1, cont).font = Font(bold=True, size=11)
        elif kind == "h2":
            ws.cell(r, 1, cont).font = BOLD
        elif kind == "p":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            c = ws.cell(r, 1, cont)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 15 * (len(cont) // 110 + 1))
        elif kind == "table":
            for j, col in enumerate(cont["cols"]):
                ws.cell(r, j + 2, col).font = BOLD
            for lab, vals in cont["rows"]:
                r += 1
                ws.cell(r, 1, lab).alignment = Alignment(indent=1)
                for j, x in enumerate(vals):
                    if x is not None:
                        ws.cell(r, j + 2, round(x, 2)).number_format = NF
            r += 1
            ws.cell(r, 1, "Total").font = BOLD
            for j, x in enumerate(cont["total"]):
                if x is not None:
                    c = ws.cell(r, j + 2, round(x, 2))
                    c.number_format = NF
                    c.font = BOLD
                    c.border = Border(top=thin)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    escribir_papeles(wb, a, X)
    escribir_pt_ganancias(wb, a, X)
    escribir_asientos(wb, a, X)

    # Control (no va al PDF)
    ws = wb.create_sheet("Control")
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 34
    ws.append(["Cuenta del sumas y saldos", "Saldo", "Rubro asignado"])
    for c in ws[1]:
        c.font = BOLD
    for n, s, k in X["cuentas"]:
        ws.append([n, round(s, 2), RUBROS[k][1]])
        ws.cell(ws.max_row, 2).number_format = NF
    ws.append([])
    ws.append(["Coeficiente de reexpresión", X["cof"]])
    ws.append(["Amortizaciones del ejercicio calculadas", round(X["amort_ej"], 2)])
    ws.append(["RFyT incl. RECPAM (por diferencia)", round(X["R"]["rfyt"], 2)])
    ws.append([])
    ws.append(["Advertencias"])
    ws.cell(ws.max_row, 1).font = BOLD
    for w in X["warn"] or ["(ninguna)"]:
        ws.append([w])
    wb.calculation.fullCalcOnLoad = True   # los papeles de trabajo usan fórmulas
    wb.save(path)


# ════════════════════════════════════════════════════════════════════════════
# Papeles de trabajo de ajuste por inflación (solo en el XLSX)
# ════════════════════════════════════════════════════════════════════════════
def escribir_papeles(wb, a, X):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    NF = '#,##0.00;(#,##0.00);"-"'
    CF = '0.000000'
    B = Font(bold=True)
    H = PatternFill("solid", fgColor="1F3864")
    INP = PatternFill("solid", fgColor="FFF2CC")   # celdas a completar / revisar
    T = PatternFill("solid", fgColor="D9E1F2")
    thin = Border(top=Side(style="thin", color="999999"))
    C, A, I, R, prior = X["C"], X["A"], X["I"], X["R"], X["prior"]
    tp = X["tiene_prior"]
    g = (lambda k: prior.g(k)) if tp else (lambda k: 0.0)
    ca, cc = X["cierre_ant"].strftime("%d/%m/%Y"), X["cierre"].strftime("%d/%m/%Y")

    def hdr(ws, r, labels):
        for j, l in enumerate(labels):
            c = ws.cell(r, j + 1, l)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = H
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30

    def titulo(ws, t):
        ws.cell(1, 1, a.empresa).font = Font(bold=True, size=12)
        ws.cell(2, 1, t).font = Font(bold=True, size=11)
        ws.cell(3, 1, f"CUIT {fmt_cuit(a.cuit)} — Ejercicio N° {a.nro_ejercicio} del "
                      f"{X['inicio'].strftime('%d/%m/%Y')} al {cc}").font = Font(italic=True, size=9)

    def num(ws, r, c, v, fmt=NF, bold=False, fill=None):
        cell = ws.cell(r, c, v)
        cell.number_format = fmt
        if bold:
            cell.font = B
        if fill:
            cell.fill = fill
        return cell

    # ───────────────────────── CONTABLE ─────────────────────────
    ws = wb.create_sheet("PT AxI contable")
    for col, w in zip("ABCDEFGHIJKLM", [44, 18, 13, 18, 18, 16, 13, 18, 18, 13, 18, 10, 18]):
        ws.column_dimensions[col].width = w
    titulo(ws, "PAPEL DE TRABAJO — AJUSTE POR INFLACIÓN CONTABLE (RT 6 / RT 54)")
    r = 5
    ws.cell(r, 1, "1. Parámetros").font = B
    r += 1
    ws.cell(r, 1, f"Coeficiente de reexpresión {ca} → {cc} (índice FACPCE / IPC INDEC)")
    num(ws, r, 2, X["cof"], CF, True, INP)
    cof_ref = f"$B${r}"
    r += 1
    ws.cell(r, 1, "Coeficiente medio del ejercicio (altas del ejercicio)")
    ws.cell(r, 2, f"=SQRT({cof_ref})").number_format = CF
    ws.cell(r, 2).fill = INP
    cofm_ref = f"$B${r}"
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=13)
    nota = ws.cell(r, 3, "Se aplica a las altas de bienes de uso cuando no se conoce el mes de compra: supone que, en "
                         "promedio, se adquirieron a mitad del ejercicio. Es la raíz cuadrada del coeficiente anual "
                         "(y no el promedio simple) porque la inflación se acumula mes a mes. Si se conoce el mes de "
                         "una compra relevante, reemplazar por el coeficiente de ese mes al cierre.")
    nota.font = Font(italic=True, size=8)
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 36
    r += 1
    ws.cell(r, 1, "Método: simplificaciones para entidades pequeñas (RT 54); resultados del ejercicio en moneda "
                  "nominal y RECPAM determinado por diferencia.").font = Font(italic=True, size=9)
    if X.get("axi_sistema"):
        r += 1
        ws.cell(r, 1, f"El sistema ya tiene registrado el ajuste por inflación (RECPAM {fmt_ar(X['recpam_ss'])}): los "
                      "estados respetan esos saldos y no se proponen asientos de reexpresión. Este papel queda "
                      "como control del cálculo con el coeficiente.").font = Font(italic=True, size=9, color="9C5700")
        if X.get("pn_control") is not None:
            r += 1
            ws.cell(r, 1, "PN de inicio: registrado en el sistema vs. PN anterior × coeficiente")
            num(ws, r, 2, X["I"]["pn"])
            num(ws, r, 3, X["pn_control"])
            num(ws, r, 4, X["I"]["pn"] - X["pn_control"], bold=True)

    # 2. PN de inicio
    r += 2
    ws.cell(r, 1, "2. Reexpresión del patrimonio neto al inicio").font = B
    r += 1
    hdr(ws, r, ["Cuenta", f"Saldo al {ca}", "Coeficiente", "Reexpresado", "Ajuste"])
    pb = X["pn_base"]
    first = r + 1
    filas_pn = [("Capital social", pb["cap"], "cap"), ("Ajuste de capital", pb["aj"], "aj"),
                ("Aportes irrevocables", pb["apirr"], None), ("Reservas", pb["resleg"], None),
                ("Resultados no asignados", pb["rna"], None)]
    for lab, v, k in filas_pn:
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v)
        ws.cell(r, 3, f"={cof_ref}").number_format = CF
        if k == "cap":
            ws.cell(r, 4, f"=B{r}").number_format = NF
            ws.cell(r, 5, "Se mantiene nominal; su ajuste va a «Ajuste de capital»").font = Font(italic=True, size=8)
        elif k == "aj":
            ws.cell(r, 4, f"=(B{r - 1}+B{r})*C{r}-B{r - 1}").number_format = NF
            ws.cell(r, 5, f"=D{r}-B{r}").number_format = NF
        else:
            ws.cell(r, 4, f"=B{r}*C{r}").number_format = NF
            ws.cell(r, 5, f"=D{r}-B{r}").number_format = NF
    r += 1
    ws.cell(r, 1, "Patrimonio neto al inicio").font = B
    for c in (2, 4, 5):
        L = "ABCDE"[c - 1]
        cell = ws.cell(r, c, f"=SUM({L}{first}:{L}{r - 1})")
        cell.number_format = NF
        cell.font = B
        cell.border = thin
    pn_ini_row = r

    # 3. Comparativos
    r += 2
    ws.cell(r, 1, "3. Reexpresión de las cifras comparativas (EECC anterior)").font = B
    r += 1
    hdr(ws, r, ["Rubro", f"Saldo al {ca}", "Coeficiente", "Reexpresado"])
    comp = [("Caja y bancos", "caja"), ("Inversiones", "inv"), ("Créditos por ventas", "cv"),
            ("Otros créditos", "oc"), ("Bienes de cambio", "bc"), ("Bienes de uso", "bu"),
            ("Deudas comerciales", "dc"), ("Préstamos", "prest"), ("Remuneraciones y cargas sociales", "rem"),
            ("Cargas fiscales", "df"), ("Deudas con socios", "ds"), ("Otras deudas", "od"),
            ("Patrimonio neto", "pn"), ("Ventas", "ventas"), ("Costo", "costo"),
            ("Gastos de comercialización", "gcom"), ("Gastos de administración", "gadm"),
            ("Resultados financieros y por tenencia", "rfyt"), ("Impuesto a las ganancias", "ig"),
            ("Resultado del ejercicio", "res")]
    for lab, k in comp:
        v = g(k)
        if k in ("costo", "gcom", "gadm", "ig"):
            v = abs(v)
        if not v:
            continue
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v)
        ws.cell(r, 3, f"={cof_ref}").number_format = CF
        ws.cell(r, 4, f"=B{r}*C{r}").number_format = NF

    # 4. Bienes de uso
    if X["bu"]:
        r += 2
        ws.cell(r, 1, "4. Bienes de uso — reexpresión y amortizaciones").font = B
        r += 1
        hdr(ws, r, ["Tipo de bien", f"Origen al {ca}", "Coeficiente", "Origen reexpresado", "Altas (nominal)",
                    "Coef. medio", "Altas reexpresadas", f"Amort. acum. al {ca}", "Amort. acum. reexpresada",
                    "Tasa %", "Amortización del ejercicio", "", "Neto al cierre"])
        b0 = r + 1
        for row in X["bu"]:
            r += 1
            ws.cell(r, 1, row["tipo"])
            num(ws, r, 2, row["o_ant"])
            ws.cell(r, 3, f"={cof_ref}").number_format = CF
            ws.cell(r, 4, f"=B{r}*C{r}").number_format = NF
            num(ws, r, 5, row["altas_nom"])
            ws.cell(r, 6, f"={cofm_ref}").number_format = CF
            ws.cell(r, 7, f"=E{r}*F{r}").number_format = NF
            num(ws, r, 8, row["a_ant"])
            ws.cell(r, 9, f"=H{r}*C{r}").number_format = NF
            num(ws, r, 10, row["tasa"], "0.00", fill=INP)
            ws.cell(r, 11, f"=MIN((D{r}+G{r})*J{r}/100,MAX(D{r}+G{r}-I{r},0))").number_format = NF
            ws.cell(r, 13, f"=D{r}+G{r}-I{r}-K{r}").number_format = NF
        r += 1
        ws.cell(r, 1, "Totales").font = B
        for L in "BDEGHIKM":
            c = ws[f"{L}{r}"]
            c.value = f"=SUM({L}{b0}:{L}{r - 1})"
            c.number_format = NF
            c.font = B
            c.border = thin
        ws.cell(r + 1, 1, "Año de alta completo. Altas reexpresadas con coeficiente medio del ejercicio "
                          "(simplificación).").font = Font(italic=True, size=8)
        r += 1

    # 5. Determinación del RECPAM
    r += 2
    ws.cell(r, 1, "5. Determinación del resultado y del RECPAM (por diferencia)").font = B
    r += 1
    hdr(ws, r, ["Concepto", "Importe"])
    det = [("Activo al cierre", C["ta"]), ("Pasivo al cierre", -C["tp"]), ("▸ Patrimonio neto al cierre", C["pn"]),
           ("(−) Patrimonio neto al inicio reexpresado", -I["pn"]), ("(−) Aportes del ejercicio", -X["aportes"]),
           ("▸ RESULTADO DEL EJERCICIO", R["res"]), ("", None),
           ("Ventas", R["ventas"]), ("(−) Costo", -R["costo"]), ("(−) Gastos de comercialización", -R["gcom"]),
           ("(−) Gastos de administración (incluye amortizaciones)", -R["gadm"]),
           ("Otros ingresos y egresos", R["oi"]), ("(−) Impuesto a las ganancias", -R["ig"]),
           ("▸ Resultado sin RFyT", R["res"] - R["rfyt"]), ("", None),
           ("RESULTADOS FINANCIEROS Y POR TENENCIA (incl. RECPAM) = Resultado − Resultado sin RFyT", R["rfyt"])]
    for lab, v in det:
        r += 1
        ws.cell(r, 1, lab)
        if v is not None:
            bold = lab.startswith("▸") or lab.startswith("RESULTADOS")
            num(ws, r, 2, v, bold=bold)
            if bold:
                ws.cell(r, 1).font = B
                ws.cell(r, 2).border = thin

    # 6. Control de razonabilidad
    r += 2
    ws.cell(r, 1, "6. Control de razonabilidad del RECPAM (posición monetaria)").font = B
    r += 1
    hdr(ws, r, ["Concepto", f"Inicio (al {ca}, nominal)", f"Cierre (al {cc})"])
    pmn_i = g("caja") + g("cv") + g("oc") - (g("dc") + g("prest") + g("rem") + g("df") + g("ds") + g("od"))
    pmn_c = C["caja"] + C["cv"] + C["oc"] - C["tp"]
    rows6 = [("Activos monetarios (caja, créditos)", g("caja") + g("cv") + g("oc"), C["caja"] + C["cv"] + C["oc"]),
             ("Pasivos monetarios", -(g("dc") + g("prest") + g("rem") + g("df") + g("ds") + g("od")), -C["tp"]),
             ("Posición monetaria neta", pmn_i, pmn_c)]
    for lab, v1, v2 in rows6:
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v1, bold=lab.startswith("Posición"))
        num(ws, r, 3, v2, bold=lab.startswith("Posición"))
    pmn_row = r
    r += 1
    ws.cell(r, 1, "RECPAM estimado sobre la posición de inicio  = −PMN inicio × (coef − 1)")
    ws.cell(r, 2, f"=-B{pmn_row}*({cof_ref}-1)").number_format = NF
    est_row = r
    r += 1
    ws.cell(r, 1, "RFyT determinado (punto 5)")
    num(ws, r, 2, R["rfyt"])
    r += 1
    ws.cell(r, 1, "Diferencia (movimientos del ejercicio, resultados nominales)").font = B
    c = ws.cell(r, 2, f"=B{r - 1}-B{est_row}")
    c.number_format = NF
    c.font = B
    c.border = thin

    # ───────────────────────── IMPOSITIVO ─────────────────────────
    ws = wb.create_sheet("PT AxI impositivo")
    for col, w in zip("ABCDEFG", [52, 18, 12, 18, 14, 18, 40]):
        ws.column_dimensions[col].width = w
    titulo(ws, "PAPEL DE TRABAJO — AJUSTE POR INFLACIÓN IMPOSITIVO (Ley de Impuesto a las Ganancias, "
               "arts. 105 a 108 t.o. 2019)")
    r = 5
    ws.cell(r, 1, "Celdas en amarillo: completar o revisar. Se parte de valores contables del EECC anterior como "
                  "aproximación de los valores impositivos.").font = Font(italic=True, size=9, color="9C5700")
    r += 2
    ws.cell(r, 1, "1. Parámetros").font = B
    r += 1
    ws.cell(r, 1, f"Variación del IPC (INDEC) del ejercicio — coeficiente {ca} → {cc}")
    num(ws, r, 2, X["cof"], CF, True, INP)
    ipc_ref = f"$B${r}"
    r += 1
    ws.cell(r, 1, "Variación acumulada del IPC en los 36 meses anteriores al cierre (%)")
    num(ws, r, 2, None, "0.00", fill=INP)
    v36 = f"B{r}"
    r += 1
    ws.cell(r, 1, "¿Procede el ajuste? (variación 36 meses > 100%)")
    ws.cell(r, 2, f'=IF({v36}="","Completar",IF({v36}>100,"SÍ","NO"))').font = B
    proc_ref = f"$B${r}"
    r += 1
    ws.cell(r, 1, "Imputación en el ejercicio (%) — 100% para ejercicios iniciados desde el 01/01/2021")
    num(ws, r, 2, 100, "0", fill=INP)
    imp_ref = f"$B${r}"

    # 2. Ajuste estático
    r += 2
    ws.cell(r, 1, f"2. Ajuste estático — activo y pasivo al inicio del ejercicio ({ca})").font = B
    r += 1
    hdr(ws, r, ["Rubro", "Saldo al inicio", "Computa (S/N)", "Importe computable", "", "", "Observación"])
    socios_ini = 0.0
    if tp:
        for n, s_, k in X["cuentas"]:
            if k == "oc" and re.search(r"socio|director|accionista", nkey(n)):
                pv = prior.cuenta(n)
                socios_ini += pv or 0.0
    act = [("Caja y bancos", g("caja"), "S", ""),
           ("Inversiones", g("inv"), "S", "Excluir acciones y participaciones en sociedades"),
           ("Créditos por ventas", g("cv"), "S", ""),
           ("Otros créditos (sin saldos de socios)", g("oc") - socios_ini, "S", ""),
           ("Saldos deudores de socios / accionistas", socios_ini, "N", "Se excluyen (art. 106 inc. a)"),
           ("Bienes de cambio", g("bc"), "S", "Revisar criterio de valuación impositiva"),
           ("Bienes de uso", g("bu"), "N", "Excluidos: se actualizan por las normas específicas"),
           ("Intangibles / gastos de organización", 0.0, "N", "Excluidos")]
    r += 1
    ws.cell(r, 1, "ACTIVO").font = B
    a0 = r + 1
    for lab, v, comp_, obs in act:
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v)
        c = ws.cell(r, 3, comp_)
        c.fill = INP
        c.alignment = Alignment(horizontal="center")
        ws.cell(r, 4, f'=IF(UPPER(C{r})="S",B{r},0)').number_format = NF
        ws.cell(r, 7, obs).font = Font(size=8, italic=True)
    r += 1
    ws.cell(r, 1, "Activo computable").font = B
    c = ws.cell(r, 4, f"=SUM(D{a0}:D{r - 1})")
    c.number_format = NF
    c.font = B
    c.border = thin
    act_row = r
    pas = [("Deudas comerciales", g("dc"), "S", ""), ("Préstamos", g("prest"), "S", ""),
           ("Remuneraciones y cargas sociales", g("rem"), "S", ""),
           ("Cargas fiscales", g("df"), "S", "Incluye la provisión del impuesto a las ganancias"),
           ("Deudas con socios", g("ds"), "S", ""), ("Otras deudas", g("od"), "S", ""),
           ("Previsiones no deducibles", 0.0, "N", "Excluidas")]
    r += 1
    ws.cell(r, 1, "PASIVO").font = B
    p0 = r + 1
    for lab, v, comp_, obs in pas:
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v)
        c = ws.cell(r, 3, comp_)
        c.fill = INP
        c.alignment = Alignment(horizontal="center")
        ws.cell(r, 4, f'=IF(UPPER(C{r})="S",B{r},0)').number_format = NF
        ws.cell(r, 7, obs).font = Font(size=8, italic=True)
    r += 1
    ws.cell(r, 1, "Pasivo computable").font = B
    c = ws.cell(r, 4, f"=SUM(D{p0}:D{r - 1})")
    c.number_format = NF
    c.font = B
    c.border = thin
    pas_row = r
    r += 1
    ws.cell(r, 1, "Diferencia (activo − pasivo computable)").font = B
    c = ws.cell(r, 4, f"=D{act_row}-D{pas_row}")
    c.number_format = NF
    c.font = B
    dif_row = r
    r += 1
    ws.cell(r, 1, "AJUSTE ESTÁTICO = −(activo − pasivo) × (IPC − 1)").font = B
    c = ws.cell(r, 4, f"=-D{dif_row}*({ipc_ref}-1)")
    c.number_format = NF
    c.font = B
    c.fill = T
    est_ref = f"D{r}"
    ws.cell(r, 7, "Negativo si activo > pasivo (quebranto); positivo si pasivo > activo").font = Font(size=8, italic=True)

    # 3. Ajuste dinámico
    r += 2
    ws.cell(r, 1, "3. Ajuste dinámico — movimientos del ejercicio").font = B
    r += 1
    hdr(ws, r, ["Concepto", "Importe", "Mes (aaaa-mm)", "", "Coef. mes → cierre", "Ajuste", "Signo / observación"])
    movs = []
    if X["aportes"]:
        movs.append(("Aportes de capital del ejercicio", X["aportes"], -1, "Aportes: ajuste negativo"))
    movs += [("Retiros / dividendos pagados", 0.0, 1, "Ajuste positivo"),
             ("Préstamos / anticipos a socios", 0.0, 1, "Ajuste positivo"),
             ("Adquisición de acciones / bienes excluidos", 0.0, 1, "Ajuste positivo"),
             ("Venta de bienes excluidos", 0.0, -1, "Ajuste negativo")]
    d0 = r + 1
    for lab, v, sg, obs in movs:
        r += 1
        ws.cell(r, 1, lab)
        num(ws, r, 2, v, fill=INP)
        ws.cell(r, 3).fill = INP
        num(ws, r, 5, None, CF, fill=INP)
        ws.cell(r, 6, f'=IF(E{r}="",0,{sg}*B{r}*(E{r}-1))').number_format = NF
        ws.cell(r, 7, obs).font = Font(size=8, italic=True)
    for _ in range(3):
        r += 1
        ws.cell(r, 1, "(otro movimiento)").font = Font(italic=True, color="999999")
        for col in (2, 3, 5):
            ws.cell(r, col).fill = INP
        ws.cell(r, 6, f'=IF(E{r}="",0,B{r}*(E{r}-1))').number_format = NF
        ws.cell(r, 7, "Cargar con signo: + positivo / − negativo").font = Font(size=8, italic=True)
    r += 1
    ws.cell(r, 1, "AJUSTE DINÁMICO").font = B
    c = ws.cell(r, 6, f"=SUM(F{d0}:F{r - 1})")
    c.number_format = NF
    c.font = B
    c.fill = T
    din_ref = f"F{r}"

    # 4. Resultado
    r += 2
    ws.cell(r, 1, "4. Resultado del ajuste por inflación impositivo").font = B
    r += 1
    ws.cell(r, 1, "Ajuste estático")
    ws.cell(r, 4, f"={est_ref}").number_format = NF
    r += 1
    ws.cell(r, 1, "Ajuste dinámico")
    ws.cell(r, 4, f"={din_ref}").number_format = NF
    r += 1
    ws.cell(r, 1, "AJUSTE POR INFLACIÓN DEL EJERCICIO").font = B
    c = ws.cell(r, 4, f"=D{r - 2}+D{r - 1}")
    c.number_format = NF
    c.font = B
    c.border = thin
    tot = f"D{r}"
    r += 1
    ws.cell(r, 1, "Imputable en el ejercicio (si procede)").font = B
    c = ws.cell(r, 4, f'=IF({proc_ref}="SÍ",{tot}*{imp_ref}/100,0)')
    c.number_format = NF
    c.font = B
    c.fill = T
    r += 1
    ws.cell(r, 1, "Positivo: se suma a la ganancia neta imponible; negativo: se resta (quebranto por inflación).") \
        .font = Font(italic=True, size=8)


# ════════════════════════════════════════════════════════════════════════════
# Liquidación del impuesto a las ganancias (provisión + cómputo Ley 25.413)
# ════════════════════════════════════════════════════════════════════════════
# Escala art. 73 LIG (sociedades), por año de inicio del ejercicio: (hasta, fijo, alícuota, excedente de)
ESCALA_IG = {
    2024: [(34703523.08, 0.0, 0.25, 0.0), (347035230.79, 8675880.77, 0.30, 34703523.08),
           (None, 102375393.08, 0.35, 347035230.79)],
    2025: [(101679575.26, 0.0, 0.25, 0.0), (1016795752.60, 25419893.82, 0.30, 101679575.26),
           (None, 299954747.02, 0.35, 1016795752.60)],
}
EXENTAS_IG = r"asociacion civil|fundacion|mutual|cooperativa"
# Cuentas del plan estándar de Xubio: aunque no figuren en el sumas y saldos (saldo cero) ya existen en el plan
PLAN_XUBIO = {nkey(x) for x in [
    "Caja", "Banco", "Valores a Depositar", "Deudores por Venta", "Otras Retenciones Sufridas",
    "Percepción de Ganancias Sufrida", "Otras Percepciones Sufridas", "Percepción Ingresos Brutos Sufrida",
    "Retención Ingresos Brutos Sufrida", "Retención Ganancias Sufrida", "Percepción de IVA Sufrida",
    "Retención de IVA Sufrida", "IVA Crédito Fiscal", "Proveedores", "Cheques Diferidos", "IVA Débito Fiscal",
    "Sueldos y Jornales a Pagar", "Cargas Sociales a pagar", "Ajuste de Capital", "Resultado del Ejercicio",
    "Venta de Servicios", "Venta de Bienes", "Compras", "Compra de Servicios", "Compra de Bienes",
    "Costo de Mercadería Vendida", "Cargas Sociales", "Sueldos y Jornales", "Resultados por Tenencia", "RECPAM",
    "Ajuste por Redondeo Decimal", "Diferencia de Cambio", "Cheques Rechazados", "Impuesto a las Ganancias",
    "Diferencia de Inventario"]}


def anio_escala(anio):
    keys = sorted(ESCALA_IG)
    return max([k for k in keys if k <= anio], default=keys[0])


def impuesto_escala(gni, anio):
    tabla = ESCALA_IG[anio_escala(anio)]
    for hasta, fijo, al, exc in tabla:
        if hasta is None or gni <= hasta:
            return round(fijo + (gni - exc) * al, 2), al
    return 0.0, 0.0


def liquidar_ganancias(a, X0):
    """Calcula la provisión de ganancias sobre el resultado del sumas y saldos (antes del cierre).
    Devuelve dict G con el detalle y G['movs'] = [(cuenta, importe_deudor, es_nueva)] para registrar."""
    G = dict(aplica=False, motivo="", movs=[], filas=[], warn=[])
    cuentas = X0["cuentas"]
    if re.search(EXENTAS_IG, nkey(a.tipo_societario or "")):
        G["motivo"] = (f"Entidad exenta del impuesto a las ganancias ({a.tipo_societario}). No se calcula provisión.")
        return G
    if abs(sum(s for _, s, k in cuentas if k == "ig")) > 0.005:
        G["motivo"] = "El sumas y saldos ya tiene registrado el impuesto a las ganancias: no se vuelve a provisionar."
        G["warn"].append(G["motivo"])
        return G
    G["aplica"] = True
    R, cof = X0["R"], X0["cof"]
    nombres = [n for n, _, _ in cuentas]

    def buscar(pat, rubro=None):
        for n, s, k in cuentas:
            if re.search(pat, nkey(n)) and (rubro is None or k == rubro):
                return n
        return None

    # 1. Resultado contable nominal (sin RECPAM): el SS viene en valores históricos
    hist = X0.get("hist")
    if hist:
        # resultado de libros (valores históricos, antes del ajuste por inflación de Xubio)
        res_keys = ("ventas", "costo", "gcom", "gadm", "amort", "oi", "rfyt")
        rfyt_nom = -sum(v for n, v in hist if clasificar(n) == "rfyt" and
                        not re.search(r"recpam|ajuste por inflacion", nkey(n)))
        res_nom = -sum(v for n, v in hist if clasificar(n) in res_keys and
                       not re.search(r"recpam|ajuste por inflacion", nkey(n)))
        if not X0["amort_booked"]:
            res_nom -= X0["amort_ej"]
        res_nom = round(res_nom, 2)
    else:
        rfyt_nom = -sum(s for n, s, k in cuentas if k == "rfyt" and not re.search(r"recpam|ajuste por inflacion", nkey(n)))
        res_nom = round(R["ventas"] - R["costo"] - R["gcom"] - R["gadm"] + R["oi"] + rfyt_nom, 2)
        if getattr(a, "ss_solo_ajustado", False):
            G["warn"].append("El sumas y saldos solo trae el saldo ajustado por inflación: la base de Ganancias se "
                             "calculó con resultados ajustados. Exportá también la columna «Saldo Final» (histórica) "
                             "para que parta de los valores de libros.")
    noded = [(n, s) for n, s, k in cuentas if k in ("gadm", "gcom", "oi") and s > 0 and
             re.search(r"multa|punitorio|recargo|donacion", nkey(n))]
    noded_tot = round(sum(s for _, s in noded), 2)

    # 2. Ajuste por inflación impositivo (arts. 105-108). Se considera que procede (IPC 36 meses > 100%).
    tp, prior = X0["tiene_prior"], X0["prior"]
    g = (lambda k: prior.g(k) or 0.0) if tp else (lambda k: 0.0)
    act = g("caja") + g("inv") + g("cv") + g("oc") + g("bc")
    pas = g("dc") + g("prest") + g("rem") + g("df") + g("ds") + g("od")
    aj_est = round(-(act - pas) * (cof - 1), 2)
    if tp:
        aportes_int, coef_ap = X0["aportes_nom"], cof ** 0.5       # aportes del ejercicio: mitad de año
    else:
        aportes_int, coef_ap = X0["aportes_nom"] - X0["integ_pend"], cof   # constitución al inicio
    aj_din = round(-aportes_int * (coef_ap - 1), 2)
    aj_imp = round((aj_est + aj_din) * (1 if a.axi_impositivo else 0), 2)
    quebr = round(a.quebrantos or 0.0, 2)
    gni = round(res_nom + noded_tot + aj_imp - quebr, 2)

    # 3. Impuesto determinado
    anio = X0["inicio"].year
    if anio not in ESCALA_IG:
        G["warn"].append(f"No tengo cargada la escala del art. 73 para ejercicios iniciados en {anio}: "
                         f"se usó la de {anio_escala(anio)}. "
                         "Verificar en el PT Ganancias.")
    imp, alic = impuesto_escala(gni, anio) if gni > 0 else (0.0, 0.25)

    # 4. Impuesto sobre débitos y créditos (Ley 25.413)
    pc = [(n, s) for n, s, k in cuentas if k == "oc" and s > 0.005 and
          re.search(r"25\.?4(13|03)|\b(db|debitos?) y (cr|creditos?)|impuesto.*debitos", nkey(n))]
    P = round(sum(s for _, s in pc), 2)
    mipyme = (a.mipyme or "").strip().lower()
    if mipyme in ("", "si", "sí", "s", "true", "1", "yes"):
        es_mipyme = True
        usado = min(P, imp)
        rem = round(P - usado, 2)
        tras = round(rem * 0.33, 2)
        gasto = round(rem - tras, 2)
    else:
        es_mipyme = False
        comp = round(P * 0.33, 2)
        usado = min(comp, imp)
        tras = round(comp - usado, 2)
        gasto = round(P - comp, 2)
    usado = round(usado, 2)

    # 5. Retenciones / percepciones / anticipos de ganancias
    rets = [(n, s) for n, s, k in cuentas if k == "oc" and s > 0.005 and
            re.search(r"(retencion|percepcion|anticipo)e?s?.*ganancia", nkey(n)) and not re.search(r"efectuad", nkey(n))]
    R_tot = round(sum(s for _, s in rets), 2)
    aplic = round(min(R_tot, max(imp - usado, 0.0)), 2)
    libre = round(R_tot - aplic, 2)
    a_pagar = round(imp - usado - aplic, 2)
    anticipos_base = round(max(imp - usado - aplic, 0.0), 2)

    # 6. Movimientos (asiento de provisión)
    ig_c = buscar(r"^impuesto a las ganancias$", "ig") or "Impuesto a las Ganancias"
    gb = buscar(r"gastos? bancari", "gadm") or buscar(r"gastos? bancari")
    gasto_c, gasto_n = (gb, False) if gb else ("Gastos bancarios", True)
    movs = [(ig_c, imp, ig_c not in nombres and nkey(ig_c) not in PLAN_XUBIO), (gasto_c, gasto, gasto_n)]
    rest = round(usado + gasto, 2)
    for n, s in pc:
        t = round(min(s, rest), 2)
        movs.append((n, -t, False))
        rest = round(rest - t, 2)
    if rets:
        for n, s in rets:
            movs.append((n, -s, False))
        if libre > 0.005:
            ld = buscar(r"libre disponibilidad.*ganancia|ganancia.*libre disponibilidad")
            movs.append((ld or "Saldo de libre disponibilidad Ganancias", libre, ld is None))
    if a_pagar > 0.005:
        ap = buscar(r"ganancias a pagar|provision.*ganancias", "df")
        movs.append((ap or "Impuesto a las ganancias a pagar", -a_pagar, ap is None))
    G["movs"] = [(c, round(v, 2), nv) for c, v, nv in movs if abs(v) > 0.005]
    G.update(res_nom=res_nom, rfyt_nom=rfyt_nom, noded=noded, noded_tot=noded_tot, aj_est=aj_est, aj_din=aj_din,
             aportes_int=aportes_int, coef_ap=coef_ap, aj_imp=aj_imp, quebr=quebr, gni=gni, anio=anio, alic=alic,
             imp=imp, P=P, pc=pc, es_mipyme=es_mipyme, usado=usado, rem=round(P - usado, 2) if es_mipyme else 0.0,
             tras=tras, gasto=gasto, rets=rets, R_tot=R_tot, aplic=aplic, libre=libre, a_pagar=a_pagar,
             anticipos_base=anticipos_base, act=act, pas=pas, cof=cof)
    return G


def aplicar_movs(cuentas, movs):
    """Suma movimientos al sumas y saldos (por nombre de cuenta); agrega las cuentas nuevas al final."""
    out = [[n, s] for n, s in cuentas]
    idx = {n: i for i, (n, _) in enumerate(out)}
    for c, v, _ in movs:
        if c in idx:
            out[idx[c]][1] = round(out[idx[c]][1] + v, 2)
        else:
            idx[c] = len(out)
            out.append([c, round(v, 2)])
    return [(n, s) for n, s in out]


def escribir_pt_ganancias(wb, a, X):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    NF = '#,##0.00;(#,##0.00);"-"'
    B = Font(bold=True)
    INP = PatternFill("solid", fgColor="FFF2CC")
    T = PatternFill("solid", fgColor="D9E1F2")
    thin = Border(top=Side(style="thin", color="999999"))
    G = X.get("ganancias") or {}
    ws = wb.create_sheet("PT Ganancias")
    for col, w in zip("ABC", [70, 20, 60]):
        ws.column_dimensions[col].width = w
    ws.cell(1, 1, a.empresa).font = Font(bold=True, size=12)
    ws.cell(2, 1, "PAPEL DE TRABAJO — IMPUESTO A LAS GANANCIAS Y CÓMPUTO LEY 25.413").font = Font(bold=True, size=11)
    ws.cell(3, 1, f"CUIT {fmt_cuit(a.cuit)} — Ejercicio N° {a.nro_ejercicio} del {X['inicio'].strftime('%d/%m/%Y')} "
                  f"al {X['cierre'].strftime('%d/%m/%Y')}").font = Font(italic=True, size=9)
    if not G.get("aplica"):
        ws.cell(5, 1, G.get("motivo") or "No se calculó la provisión.").font = B
        return
    r = 4

    def sec(t):
        nonlocal r
        r += 2
        ws.cell(r, 1, t).font = B

    def fila(lab, v, bold=False, fill=None, obs="", fmt=NF):
        nonlocal r
        r += 1
        ws.cell(r, 1, lab).font = B if bold else Font()
        c = ws.cell(r, 2, v)
        c.number_format = fmt
        if bold:
            c.font = B
            c.border = thin
        if fill:
            c.fill = fill
        if obs:
            ws.cell(r, 3, obs).font = Font(italic=True, size=8)
        return r

    sec("1. Determinación de la ganancia neta imponible")
    fila("Resultado contable en valores históricos (sin RECPAM)", G["res_nom"],
         obs="Ventas − costo − gastos (incl. amortizaciones) + otros ingresos + resultados financieros nominales")
    for n, s in G["noded"]:
        fila(f"(+) No deducible: {n}", s, fill=INP, obs="Revisar deducibilidad")
    fila("(+/−) Ajuste por inflación impositivo (ver PT AxI impositivo / punto 2)", G["aj_imp"])
    fila("(−) Quebrantos de ejercicios anteriores", -G["quebr"], fill=INP)
    fila("▸ GANANCIA NETA IMPONIBLE", G["gni"], bold=True)
    fila("Alícuota marginal (art. 73, escala del ejercicio iniciado en " + str(G["anio"]) + ")", G["alic"],
         fill=INP, fmt="0%", obs="Verificar escala vigente publicada por ARCA")
    fila("▸ IMPUESTO DETERMINADO", G["imp"], bold=True, fill=T)

    sec("2. Ajuste por inflación impositivo (cálculo usado en la provisión)")
    fila("Activo computable al inicio (EECC anterior)", G["act"])
    fila("Pasivo computable al inicio (EECC anterior)", G["pas"])
    fila("Ajuste estático = −(activo − pasivo) × (coef − 1)", G["aj_est"])
    fila("Aportes de capital integrados del ejercicio", G["aportes_int"],
         obs="Primer ejercicio: capital integrado (sin el pendiente de integración)" if not X["tiene_prior"] else "")
    fila("Coeficiente aplicado a los aportes", G["coef_ap"], fill=INP, fmt="0.000000",
         obs="Reemplazar por el coeficiente IPC del mes del aporte (tabla ARCA)")
    fila("Ajuste dinámico = −aportes × (coef − 1)", G["aj_din"])
    fila("▸ AJUSTE POR INFLACIÓN IMPOSITIVO", G["aj_imp"], bold=True,
         obs="Se asume que procede (variación IPC 36 meses > 100%), imputación 100%")

    sec("3. Impuesto sobre los débitos y créditos bancarios (Ley 25.413)")
    fila("Pago a cuenta disponible", G["P"])
    fila("Condición", "MiPyME (100% computable)" if G["es_mipyme"] else "No MiPyME (33% computable)",
         fmt="@", fill=INP)
    fila("Computado contra el impuesto determinado", G["usado"])
    if G["es_mipyme"]:
        fila("Remanente", G["rem"])
    fila("▸ Saldo trasladable al ejercicio siguiente (33%)", G["tras"], bold=True)
    fila("▸ Imputado a gasto (no trasladable)", G["gasto"], bold=True)

    sec("4. Retenciones, percepciones y anticipos de ganancias")
    for n, s in G["rets"]:
        fila(n, s)
    fila("Aplicados contra el impuesto", G["aplic"])
    fila("▸ Saldo de libre disponibilidad", G["libre"], bold=True)

    sec("5. Saldo de la declaración jurada")
    fila("Impuesto determinado", G["imp"])
    fila("(−) Ley 25.413 computada", -G["usado"])
    fila("(−) Retenciones / percepciones / anticipos", -G["aplic"])
    fila("▸ IMPUESTO A INGRESAR", G["a_pagar"], bold=True, fill=T)
    fila("Base para anticipos del ejercicio siguiente", G["anticipos_base"])

    sec("6. Asiento de provisión (incluido en «Asientos de ajuste» y en el archivo de importación)")
    for c, v, nv in G["movs"]:
        r += 1
        ws.cell(r, 1, ("" if v > 0 else "      a ") + c + ("  (cuenta nueva)" if nv else ""))
        ws.cell(r, 2 if v > 0 else 3, abs(v)).number_format = NF


# ════════════════════════════════════════════════════════════════════════════
# Asientos para llevar el sistema al balance presentado
# ════════════════════════════════════════════════════════════════════════════
def armar_asientos(a, X):
    """Devuelve (asientos, control). Cada asiento: dict(n, titulo, lineas=[(cuenta, debe, haber, nueva)])."""
    cuentas = X["cuentas"]
    nombres = {nkey(n): n for n, _, _ in cuentas}
    # saldo inicial = sumas y saldos tal como vino del sistema (sin la provisión de ganancias)
    ss = X.get("cuentas_ss") or [(n, s) for n, s, _ in cuentas]
    libros = X.get("hist")                          # saldos históricos (de libros), si el SS los trae
    en_ss = {n for n, _ in ss}
    saldo = {n: s for n, s in (libros or ss)}       # saldo deudor positivo

    def cta(preferidas, nueva):
        for p in preferidas:
            for kn, n in nombres.items():
                if re.search(p, kn):
                    return n, False
        return nueva, nkey(nueva) not in PLAN_XUBIO

    asientos = []

    def asiento(titulo, movs):
        """movs: lista de (cuenta, importe_deudor, es_nueva); positivo = Debe."""
        lineas = []
        for c, imp, nueva in movs:
            imp = round(imp, 2)
            if abs(imp) < 0.005:
                continue
            lineas.append((c, imp if imp > 0 else 0.0, -imp if imp < 0 else 0.0, nueva))
            saldo[c] = saldo.get(c, 0.0) + imp
        if lineas:
            asientos.append(dict(n=len(asientos) + 1, titulo=titulo, lineas=lineas))

    recpam, recpam_n = cta([r"^recpam$", r"recpam"], "RECPAM")
    rna_c, rna_n = cta([r"resultados? no asignados?", r"^rna$"], "Resultados no asignados")

    # 0. Ajuste por inflación calculado por Xubio (diferencia entre saldo ajustado y saldo de libros)
    if libros:
        adj = {n: s for n, s in ss}
        movs = [(n, round(adj.get(n, 0.0) - saldo.get(n, 0.0), 2), False) for n in adj]
        asiento("Ajuste por inflación según Xubio (saldo ajustado − saldo histórico). No importar si ya está "
                "registrado en el sistema", movs)

    # 1. Reclasificación del resultado del ejercicio anterior a Resultados no asignados
    movs = []
    for n, s, k in cuentas:
        if k == "rna" and n != rna_c and abs(s) > 0.005:
            movs += [(n, -s, False), (rna_c, s, rna_n)]
    asiento("Reclasificación del resultado del ejercicio anterior a Resultados no asignados", movs)

    # 2. Reexpresión del PN de inicio
    if X["tiene_prior"] and not X["axi_sistema"]:
        I, sp = X["I"], X["ss_pn"]
        aj_c, aj_n = cta([r"ajuste (de|del) capital"], "Ajuste de capital")
        movs = [(aj_c, -(I["aj"] - sp["aj"]), aj_n)]
        if abs(I["apirr"] - sp["apirr"]) > 0.005:
            c, nv = cta([r"ajuste.*aportes"], "Ajuste de aportes irrevocables")
            movs.append((c, -(I["apirr"] - sp["apirr"]), nv))
        if abs(I["resleg"] - sp["resleg"]) > 0.005:
            c, nv = cta([r"^reserva"], "Reserva legal")
            movs.append((c, -(I["resleg"] - sp["resleg"]), nv))
        movs.append((rna_c, -(I["rna"] - sp["rna"]), rna_n))
        movs.append((recpam, -sum(m[1] for m in movs), recpam_n))
        asiento(f"Reexpresión del patrimonio neto al inicio (coeficiente {X['cof']:.5f}".replace(".", ",") + ")", movs)
    elif not X["tiene_prior"] and X.get("aportes_aj") and not X["axi_sistema"]:
        aj_c, aj_n = cta([r"ajuste (de|del) capital"], "Ajuste de Capital")
        falta = round(X["aportes_aj"] - X["ss_pn"]["aj"], 2)     # lo que el sistema todavía no registró
        asiento(f"Reexpresión del capital suscripto en el ejercicio (coeficiente {X['cof']:.5f}".replace(".", ",") + ")",
                [(recpam, falta, recpam_n), (aj_c, -falta, aj_n)])

    # 3. Reexpresión de bienes de uso y 4. amortizaciones del ejercicio
    movs3, movs4 = [], []
    gasto_c, gasto_n = cta([r"^amortizacion(es)? (de )?bienes", r"^amortizacion(?!.*acum)"], "Amortizaciones bienes de uso")
    for row in X["bu"]:
        orig = [(n, s) for n, s, k in cuentas if k == "bu" and tipo_bu(n) == row["tipo"] and
                not re.search(r"amort|depreciacion", nkey(n))]
        acum = [(n, s) for n, s, k in cuentas if k == "bu" and tipo_bu(n) == row["tipo"] and
                re.search(r"amort|depreciacion", nkey(n))]
        o_c, o_n = (orig[0][0], False) if orig else (row["tipo"], True)
        a_c, a_n = (acum[0][0], False) if acum else (f"Amortización acumulada {row['tipo'].lower()}", True)
        o_ss = sum(v for _, v in orig)
        a_ss = sum(v for _, v in acum)            # saldo acreedor → negativo
        d_orig = row["o_cie"] - o_ss
        d_acum = -row["a_ini"] - a_ss
        if X["amort_booked"]:
            d_acum = -(row["a_cie"]) - a_ss
        movs3 += [(o_c, d_orig, o_n), (a_c, d_acum, a_n)]
        if not X["amort_booked"] and row["amort"]:
            movs4 += [(gasto_c, row["amort"], gasto_n), (a_c, -row["amort"], a_n)]
    if movs3:
        movs3.append((recpam, -sum(m[1] for m in movs3), recpam_n))
        asiento("Reexpresión de bienes de uso y de sus amortizaciones acumuladas al inicio", movs3)
    asiento("Amortizaciones de bienes de uso del ejercicio", movs4)

    # Provisión del impuesto a las ganancias y cómputo de la Ley 25.413
    G = X.get("ganancias") or {}
    if G.get("aplica") and G.get("movs"):
        asiento("Provisión del impuesto a las ganancias, cómputo Ley 25.413 y retenciones (ver PT Ganancias)",
                [(c, v, nv) for c, v, nv in G["movs"]])

    # 5. Refundición de cuentas de resultado
    res_c, res_n = cta([r"^resultados? del ejercicio$"], "Resultado del Ejercicio")
    movs, tot = [], 0.0
    for c, v in list(saldo.items()):
        k = clasificar(c)
        if k in ("ventas", "costo", "gcom", "gadm", "amort", "rfyt", "oi", "ig") and abs(v) > 0.005:
            movs.append((c, -v, c not in en_ss and nkey(c) not in PLAN_XUBIO))
            tot += v
    if movs:
        movs.append((res_c, tot, res_n))
        asiento("Refundición de cuentas de resultado (si el sistema no la hace automáticamente)", movs)

    # Control: sistema después de los asientos vs. balance presentado
    rub = {}
    for c, v in saldo.items():
        k = clasificar(c)
        if k == "caja" and v < -0.5:
            k = "prest"
        rub[k] = rub.get(k, 0.0) + v
    C = X["C"]
    ctrl = []
    for k, lab, sg in [("caja", "Caja y bancos", 1), ("inv", "Inversiones", 1), ("cv", "Créditos por ventas", 1),
                       ("oc", "Otros créditos", 1), ("bc", "Bienes de cambio", 1), ("bu", "Bienes de uso", 1),
                       ("dc", "Deudas comerciales", -1), ("prest", "Préstamos", -1),
                       ("rem", "Remuneraciones y cargas sociales", -1), ("df", "Cargas fiscales", -1),
                       ("ds", "Deudas con socios", -1), ("od", "Otras deudas", -1)]:
        sis = sg * rub.get(k, 0.0)
        if abs(sis) > 0.005 or abs(C.get(k, 0)) > 0.005:
            ctrl.append((lab, sis, C.get(k, 0.0)))
    pn_sis = -sum(rub.get(k, 0.0) for k in ("cap", "aj", "apirr", "resleg", "rna"))
    ctrl.append(("Patrimonio neto", pn_sis, C["pn"]))
    return asientos, ctrl


def escribir_asientos(wb, a, X):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    NF = '#,##0.00;(#,##0.00);"-"'
    B = Font(bold=True)
    H = PatternFill("solid", fgColor="1F3864")
    NEW = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(top=Side(style="thin", color="999999"))
    asientos, ctrl = armar_asientos(a, X)
    X["asientos"] = asientos
    ws = wb.create_sheet("Asientos de ajuste")
    for col, w in zip("ABCDEF", [8, 12, 52, 18, 18, 30]):
        ws.column_dimensions[col].width = w
    ws.cell(1, 1, a.empresa).font = Font(bold=True, size=12)
    ws.cell(2, 1, "ASIENTOS A REGISTRAR PARA QUE EL SISTEMA COINCIDA CON EL BALANCE PRESENTADO").font = Font(bold=True,
                                                                                                         size=11)
    ws.cell(3, 1, "Parten del sumas y saldos adjuntado (antes del cierre, en valores históricos). Cuentas en amarillo: "
                  "no existen en el sumas y saldos, hay que darlas de alta en el plan de cuentas.") \
        .font = Font(italic=True, size=9)
    r = 5
    fecha = X["cierre"].strftime("%d/%m/%Y")
    for asi in asientos:
        r += 1
        ws.cell(r, 1, f"Asiento {asi['n']}").font = B
        ws.cell(r, 3, asi["titulo"]).font = B
        r += 1
        for j, h in enumerate(["N°", "Fecha", "Cuenta", "Debe", "Haber", ""]):
            c = ws.cell(r, j + 1, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = H
        r0 = r + 1
        for cta_, d, h, nueva in sorted(asi["lineas"], key=lambda l: l[1] == 0):
            r += 1
            ws.cell(r, 1, asi["n"])
            ws.cell(r, 2, fecha)
            c = ws.cell(r, 3, ("      a " if not d else "") + cta_)
            if nueva:
                c.fill = NEW
                ws.cell(r, 6, "cuenta nueva").font = Font(italic=True, size=8)
            if d:
                ws.cell(r, 4, round(d, 2)).number_format = NF
            if h:
                ws.cell(r, 5, round(h, 2)).number_format = NF
        r += 1
        ws.cell(r, 3, "Totales").font = B
        for col in (4, 5):
            L = "D" if col == 4 else "E"
            c = ws.cell(r, col, f"=SUM({L}{r0}:{L}{r - 1})")
            c.number_format = NF
            c.font = B
            c.border = thin
        ws.cell(r, 6, f'=IF(ROUND(D{r}-E{r},2)=0,"✓ balanceado","✗ revisar")').font = B
        r += 1
    if not asientos:
        r += 1
        ws.cell(r, 1, "No hacen falta asientos: el sumas y saldos ya coincide con el balance.")
    if X["R"]["usa_inv"]:
        r += 2
        ws.cell(r, 1, "Bienes de cambio: el costo se calculó por diferencia de inventario. Registrar el asiento de "
                      "costo de ventas según el criterio del sistema (inventario periódico o permanente).") \
            .font = Font(italic=True, color="9C5700")
    r += 2
    ws.cell(r, 1, "Control: saldos del sistema después de los asientos vs. balance presentado").font = B
    r += 1
    for j, h in enumerate(["", "", "Rubro", "Sistema", "Balance", "Diferencia"]):
        c = ws.cell(r, j + 1, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = H
    for lab, sis, bal in ctrl:
        r += 1
        ws.cell(r, 3, lab)
        ws.cell(r, 4, round(sis, 2)).number_format = NF
        ws.cell(r, 5, round(bal, 2)).number_format = NF
        ws.cell(r, 6, f"=D{r}-E{r}").number_format = NF


def escribir_xubio(path, a, X):
    """Archivo de importación de asientos manuales de Xubio (modelo FECHA | CONCEPTO | CIRCUITOCONTABLE | CUENTA |
    DEBE | HABER | ORGANIZACION | CENTRODECOSTO | DESCRIPCION): una fila de cabecera por asiento y luego sus líneas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    cols = ["FECHA", "CONCEPTO", "CIRCUITOCONTABLE", "CUENTA", "DEBE", "HABER", "ORGANIZACION", "CENTRODECOSTO",
            "DESCRIPCION"]
    ws.append(cols)
    for c, w in zip("ABCDEFGHI", [12, 70, 18, 42, 16, 16, 18, 16, 40]):
        ws.column_dimensions[c].width = w
    NEW = PatternFill("solid", fgColor="FFF2CC")
    fecha = datetime.combine(X["cierre"], datetime.min.time())
    for asi in X.get("asientos", []):
        ws.append([fecha, f"Ajuste cierre {asi['n']} - {asi['titulo']}"[:120], "default"])
        ws.cell(ws.max_row, 1).number_format = "dd/mm/yyyy"
        for cta_, d, h, nueva in sorted(asi["lineas"], key=lambda l: l[1] == 0):
            ws.append([None, None, None, cta_, round(d, 2) if d else None, round(h, 2) if h else None, None, None,
                       "Cuenta nueva: dar de alta en el plan antes de importar" if nueva else None])
            if nueva:
                ws.cell(ws.max_row, 4).fill = NEW
    wb.save(path)


def subtitulo(a, X):
    return (f"Ejercicio N° {a.nro_ejercicio} iniciado el {X['inicio'].strftime('%d/%m/%Y')} y "
            f"finalizado el {X['cierre'].strftime('%d/%m/%Y')}" +
            (" — comparativo con el ejercicio anterior" if X["tiene_prior"] else "") +
            " — en moneda homogénea")


def caratula_lineas(a, X):
    return [
        ("Denominación", a.empresa),
        ("CUIT", fmt_cuit(a.cuit)),
        ("Domicilio legal", a.domicilio or ""),
        ("Inscripción IGJ", a.matricula_igj or ""),
        ("Actividad principal", X["actividad"] or ""),
        ("Ejercicio económico N°", str(a.nro_ejercicio)),
        ("Iniciado el", fecha_larga(X["inicio"])),
        ("Finalizado el", fecha_larga(X["cierre"])),
        ("Capital suscripto e integrado", f"$ {fmt_ar(a.cap_nominal or X['I']['cap'] + X['aportes'])}"),
        ("Normas contables", "RT 54 (T.O. RT 59) — FACPCE / RP N° 460/2024 CPCECABA"),
        ("Unidad de medida", f"Moneda homogénea al {X['cierre'].strftime('%d/%m/%Y')}"),
    ]


# ════════════════════════════════════════════════════════════════════════════
# Salida PDF (ReportLab)
# ════════════════════════════════════════════════════════════════════════════
def hoy_ar():
    """Fecha de hoy en Buenos Aires (el servidor corre en UTC)."""
    from zoneinfo import ZoneInfo
    return datetime.now(ZoneInfo("America/Argentina/Buenos_Aires")).date()


def cargo_por_tipo(tipo, empresa):
    t = nkey(tipo or "")
    e = nkey(empresa or "")
    if "srl" in t or re.search(r"\bs\.?\s?r\.?\s?l\.?$", e):
        return "Socio Gerente"
    if "sas" in t or re.search(r"\bs\.?\s?a\.?\s?s\.?$", e):
        return "Administrador Titular"
    return "Presidente"


def escribir_pdf(path, a, X, S, notas):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table,
                                    TableStyle, NextPageTemplate, PageBreak, KeepTogether)

    AZUL = colors.HexColor("#1F3864")
    TOTF = colors.HexColor("#E7ECF6")
    GRF = colors.HexColor("#C9D5EC")
    st_t = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=12, leading=15, alignment=TA_CENTER)
    st_s = ParagraphStyle("s", fontName="Helvetica-Oblique", fontSize=8, leading=10, alignment=TA_CENTER)
    st_h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                           spaceAfter=3)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=9, leading=12, spaceBefore=5, spaceAfter=2)
    st_p = ParagraphStyle("p", fontName="Helvetica", fontSize=9, leading=12, alignment=TA_JUSTIFY, spaceAfter=3)
    st_c = ParagraphStyle("c", fontName="Helvetica", fontSize=7.5, leading=9)
    st_cb = ParagraphStyle("cb", parent=st_c, fontName="Helvetica-Bold")
    st_hd = ParagraphStyle("hd", fontName="Helvetica-Bold", fontSize=7, leading=8.5, alignment=TA_CENTER,
                           textColor=colors.white)

    import os
    from reportlab.lib.utils import ImageReader
    firma_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "firma_contador.png")
    firma = ImageReader(firma_path) if os.path.exists(firma_path) else None
    cargo = a.cargo or cargo_por_tipo(a.tipo_societario, a.empresa)
    f_inf = (datetime.strptime(a.fecha_informe, "%Y-%m-%d").date() if a.fecha_informe else hoy_ar())
    leyenda = f"Firmado a efectos de su identificación con mi informe de fecha {f_inf.strftime('%d/%m/%Y')}"
    FIRMA_H = 0.0         # las firmas van a continuación de cada cuadro

    def bloque_firmas(width):
        """Leyenda + firma del representante (izq.) y firma/sello del contador (der.), como flowable."""
        from reportlab.platypus import Image as RLImage
        st_n = ParagraphStyle("fn", fontName="Helvetica-Bold", fontSize=8, leading=10, alignment=TA_CENTER)
        st_c2 = ParagraphStyle("fc", fontName="Helvetica", fontSize=7.5, leading=9.5, alignment=TA_CENTER)
        izq = Table([[""], [Paragraph(a.representante or "&nbsp;", st_n)], [Paragraph(cargo, st_c2)],
                     [Paragraph(a.empresa, st_c2)]], colWidths=[6.4 * cm], rowHeights=[2.0 * cm, None, None, None])
        izq.setStyle(TableStyle([("LINEBELOW", (0, 0), (0, 0), 0.5, colors.black),
                                 ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        if firma:
            iw, ih = firma.getSize()
            fh = 2.9 * cm
            der = RLImage(firma_path, width=fh * iw / ih, height=fh)
        else:
            der = ""
        width = min(width, 17.0 * cm)   # en hojas apaisadas, el bloque queda centrado y más junto
        t = Table([[izq, "", der]], colWidths=[6.6 * cm, width - 13.2 * cm, 6.6 * cm], hAlign="CENTER")
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("ALIGN", (2, 0), (2, 0), "CENTER"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        return t

    def deco(canv, doc):
        canv.saveState()
        w, h = canv._pagesize
        canv.setFont("Helvetica-Bold", 8)
        canv.drawString(1.5 * cm, h - 1.0 * cm, a.empresa)
        canv.setFont("Helvetica", 7.5)
        canv.drawRightString(w - 1.5 * cm, h - 1.0 * cm, f"CUIT {fmt_cuit(a.cuit)}"
                             + (f" — Inscripción IGJ N° {a.matricula_igj}" if a.matricula_igj else ""))
        canv.setStrokeColor(AZUL)
        canv.line(1.5 * cm, h - 1.2 * cm, w - 1.5 * cm, h - 1.2 * cm)
        canv.setFont("Helvetica", 6.5)
        canv.drawRightString(w - 1.5 * cm, 0.45 * cm, f"Página {doc.page}")
        canv.restoreState()

    def deco_caratula(canv, doc):
        canv.saveState()
        w, h = canv._pagesize
        canv.setFont("Helvetica", 6.5)
        canv.drawRightString(w - 1.5 * cm, 0.45 * cm, f"Página {doc.page}")
        canv.restoreState()

    pw, ph = A4
    lw, lh = landscape(A4)
    doc = BaseDocTemplate(path, pagesize=A4, title=f"EECC {a.empresa} {X['cierre'].year}",
                          author="Juan Pablo Stivala")
    fr_c = Frame(1.5 * cm, 1.4 * cm, pw - 3 * cm, ph - 3.0 * cm, id="c")
    fr_p = Frame(1.5 * cm, 1.0 * cm, pw - 3 * cm, ph - 2.6 * cm, id="p")
    fr_l = Frame(1.5 * cm, 1.0 * cm, lw - 3 * cm, lh - 2.6 * cm, id="l")
    doc.addPageTemplates([PageTemplate(id="C", frames=[fr_c], pagesize=A4, onPage=deco_caratula),
                          PageTemplate(id="P", frames=[fr_p], pagesize=A4, onPage=deco),
                          PageTemplate(id="L", frames=[fr_l], pagesize=landscape(A4), onPage=deco)])
    E = []

    # Carátula
    E.append(Spacer(1, 3 * cm))
    E.append(Paragraph("DENOMINACIÓN DE LA ENTIDAD", st_s))
    E.append(Paragraph(a.empresa, ParagraphStyle("big", parent=st_t, fontSize=18, leading=24)))
    E.append(Spacer(1, 0.4 * cm))
    E.append(Paragraph(f"ESTADOS CONTABLES CORRESPONDIENTES AL EJERCICIO ECONÓMICO N° {a.nro_ejercicio}", st_t))
    E.append(Paragraph(f"iniciado el {fecha_larga(X['inicio'])} y finalizado el {fecha_larga(X['cierre'])}",
                       ParagraphStyle("s2", parent=st_s, fontSize=10, leading=13)))
    E.append(Paragraph("(Expresados en moneda homogénea — pesos — a la fecha de cierre)", st_s))
    E.append(Spacer(1, 1.2 * cm))
    data = [[Paragraph(k, st_cb), Paragraph(v or "", st_c)] for k, v in caratula_lineas(a, X)]
    t = Table(data, colWidths=[5.5 * cm, 11 * cm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#BBBBBB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    E.append(t)
    cur = "C"

    for stt in S:
        tpl = "L" if stt["land"] else "P"
        if tpl != cur:
            E.append(NextPageTemplate(tpl))
            cur = tpl
        E.append(PageBreak())
        E.append(Paragraph(stt["title"], st_t))
        E.append(Paragraph(subtitulo(a, X), st_s))
        E.append(Spacer(1, 0.35 * cm))
        width = (lw if stt["land"] else pw) - 3 * cm
        n = len(stt["cols"])
        if stt["land"]:
            numw = min(3.4 * cm, (width - 6.5 * cm) / n)
        else:
            numw = 2.6 * cm if n <= 2 else min(2.6 * cm, (width - 6.0 * cm) / n)
        labw = width - numw * n
        data = [[Paragraph("", st_hd)] + [Paragraph(c, st_hd) for c in stt["cols"]]]
        styles = [("BACKGROUND", (0, 0), (-1, 0), AZUL), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                  ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 1), (-1, -1), 7.5),
                  ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
        for i, row in enumerate(stt["rows"], 1):
            kind, lab, vals = row[0], row[1], row[2]
            txt = row[3] if len(row) > 3 else {}
            bold = kind in ("sec", "sub", "tot", "grand")
            ind = "&nbsp;&nbsp;&nbsp;" if kind == "det" else ""
            cells = [Paragraph(ind + lab, st_cb if bold else st_c)]
            for j in range(n):
                x = (vals or [None] * n)[j] if vals and j < len(vals) else None
                s = txt[j] if j in txt else (fmt_ar(x) if x is not None else "")
                cells.append(s)
            data.append(cells)
            if bold:
                styles.append(("FONTNAME", (1, i), (-1, i), "Helvetica-Bold"))
            if kind == "tot":
                styles += [("BACKGROUND", (0, i), (-1, i), TOTF), ("LINEABOVE", (1, i), (-1, i), 0.5, colors.grey)]
            if kind == "grand":
                styles += [("BACKGROUND", (0, i), (-1, i), GRF), ("LINEABOVE", (1, i), (-1, i), 0.8, AZUL)]
            for j in txt:
                styles.append(("ALIGN", (j + 1, i), (j + 1, i), "CENTER"))
        t = Table(data, colWidths=[labw] + [numw] * n, repeatRows=1)
        t.setStyle(TableStyle(styles))
        E.append(t)
        E.append(KeepTogether([
            Spacer(1, 0.2 * cm),
            Paragraph("Las notas y anexos que se acompañan forman parte integrante de los estados contables.",
                      ParagraphStyle("ley", fontName="Helvetica-Oblique", fontSize=7, leading=9)),
            Spacer(1, 0.5 * cm),
            bloque_firmas(width)]))

    # Notas
    if cur != "P":
        E.append(NextPageTemplate("P"))
    E.append(PageBreak())
    E.append(Paragraph("NOTAS A LOS ESTADOS CONTABLES", st_t))
    E.append(Paragraph(subtitulo(a, X), st_s))
    E.append(Spacer(1, 0.3 * cm))
    for kind, cont in notas:
        if kind == "h1":
            E.append(Paragraph(cont, st_h1))
        elif kind == "h2":
            E.append(Paragraph(cont, st_h2))
        elif kind == "p":
            E.append(Paragraph(cont, st_p))
        elif kind == "table":
            n = len(cont["cols"])
            data = [[""] + [Paragraph(c, ParagraphStyle("nh", parent=st_cb, alignment=2)) for c in cont["cols"]]]
            for lab, vals in cont["rows"]:
                data.append([Paragraph(lab, st_c)] + [fmt_ar(x) if x is not None else "" for x in vals])
            data.append([Paragraph("Total", st_cb)] + [fmt_ar(x) if x is not None else "" for x in cont["total"]])
            t = Table(data, colWidths=[11.5 * cm - (n - 1) * 2.8 * cm + 2.8 * cm] + [2.8 * cm] * n)
            t.setStyle(TableStyle([("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                                   ("FONTNAME", (1, -1), (-1, -1), "Helvetica-Bold"),
                                   ("LINEABOVE", (1, -1), (-1, -1), 0.5, colors.grey),
                                   ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
            E.append(KeepTogether([t, Spacer(1, 0.2 * cm)]))
    doc.build(E)


# ════════════════════════════════════════════════════════════════════════════
# Main
# ════════════════════════════════════════════════════════════════════════════
def main(argv=None):
    ap = argparse.ArgumentParser(description="Genera EECC (RT 54) desde sumas y saldos + EECC anterior")
    ap.add_argument("--empresa", required=True)
    ap.add_argument("--cuit", required=True)
    ap.add_argument("--nro-ejercicio", type=int, default=1)
    ap.add_argument("--fecha-cierre", required=True, help="YYYY-MM-DD")
    ap.add_argument("--fecha-inicio", default="", help="YYYY-MM-DD (opcional, primer ejercicio)")
    ap.add_argument("--cof", type=float, required=True)
    ap.add_argument("--cap-nominal", type=float, default=0.0)
    ap.add_argument("--ss-actual", required=True)
    ap.add_argument("--eecc-anterior", default=None)
    ap.add_argument("--domicilio", default="")
    ap.add_argument("--matricula-igj", default="")
    ap.add_argument("--actividad", default="")
    ap.add_argument("--representante", default="", help="Nombre y apellido del representante legal")
    ap.add_argument("--cargo", default="", help="Cargo del representante (si se omite, según tipo societario)")
    ap.add_argument("--tipo-societario", default="")
    ap.add_argument("--fecha-informe", default="", help="YYYY-MM-DD (por defecto, hoy)")
    ap.add_argument("--mipyme", default="", help="Sí/No: impuesto a los débitos y créditos computable al 100%%")
    ap.add_argument("--quebrantos", type=float, default=0.0, help="Quebrantos impositivos de ejercicios anteriores")
    ap.add_argument("--sin-axi-impositivo", dest="axi_impositivo", action="store_false",
                    help="No computar el ajuste por inflación impositivo")
    ap.add_argument("--xubio", default=None, help="ruta del .xlsx de asientos para importar en Xubio")
    ap.add_argument("--output", required=True, help="ruta del .xlsx")
    ap.add_argument("--pdf", default=None, help="ruta del .pdf (por defecto junto al xlsx)")
    a = ap.parse_args(argv)
    a.empresa = a.empresa.strip()
    try:
        if a.cof <= 0:
            raise EECCError("El coeficiente de reexpresión tiene que ser mayor a cero.")
        if a.nro_ejercicio > 1 and not a.eecc_anterior:
            raise EECCError("Falta el EECC del ejercicio anterior (PDF) para los comparativos.")
        cuentas, ss_info = leer_ss(a.ss_actual)
        a.ss_ajustado = bool(ss_info["ajustado"])
        a.ss_solo_ajustado = bool(ss_info["solo_ajustado"])
        prior = Prior(a.eecc_anterior if a.nro_ejercicio > 1 else None)
        if a.nro_ejercicio > 1 and not prior.text.strip():
            raise EECCError("No pude leer el texto del EECC anterior (¿PDF escaneado?).")
        if a.nro_ejercicio > 1 and prior.v.get("pn") is None:
            raise EECCError("No encontré el Patrimonio Neto en el EECC anterior: revisar que el PDF sea el balance "
                            "del ejercicio anterior de esta sociedad.")
        cuit_pdf = re.sub(r"\D", "", prior.text[:4000])
        X = calcular(a, cuentas, prior)
        X["hist"] = ss_info["hist"]
        G = liquidar_ganancias(a, X)
        if G["aplica"] and G["movs"]:
            X = calcular(a, aplicar_movs(cuentas, G["movs"]), prior)
        X["ganancias"] = G
        X["cuentas_ss"] = cuentas
        X["hist"] = ss_info["hist"]
        if a.ss_solo_ajustado:
            X["warn"].append("El sumas y saldos solo trae el saldo ajustado por inflación: los asientos parten de "
                             "esos saldos (como si el ajuste ya estuviera registrado en Xubio). Para que partan de los "
                             "saldos de libros, exportá también la columna «Saldo Final».")
        X["warn"] += G["warn"]
        S, notas, anx, nota_n = armar(a, X)
        escribir_xlsx(a.output, a, X, S, notas)
        xubio_path = a.xubio or re.sub(r"\.xlsx$", "", a.output) + "_xubio.xlsx"
        escribir_xubio(xubio_path, a, X)
        pdf_path = a.pdf or re.sub(r"\.xlsx$", "", a.output) + ".pdf"
        escribir_pdf(pdf_path, a, X, S, notas)
        meta = dict(notas_rango=f"1 a {4 if X['C']['pn'] < 0 else 3}",
                    anexos_rango=("I" if len(anx) == 1 else f"I a {list(anx.values())[-1]}"),
                    advertencias=X["warn"], pdf=pdf_path, xubio=xubio_path)
        with open(re.sub(r"\.xlsx$", "", a.output) + ".meta.json", "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)
        for w in X["warn"]:
            print("ADVERTENCIA:", w)
        print(f"OK {a.output} {pdf_path}")
        return 0
    except EECCError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
